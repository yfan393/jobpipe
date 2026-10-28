#!/usr/bin/env python3
"""jobpipe v2: the mechanical half of Yuanting's unified daily job search, so Claude spends its
effort on the browser-only sites, on reading postings and on the final score.

Stdlib only. Runs in the Claude workspace, on GitHub Actions, or on a laptop. All rules live in
config.json; companies live in references/companies.csv.

The unified daily run, in order (references/unified-run.md has the full procedure):
  doctor            what this machine can reach (job boards direct? GitHub? HN?)
  collect           every employer feed + SimplifyJobs + speedyapply + Ask HN -> raw.json      [network]
  plan              the ledger: every API source, every (site x query family) and every career page
  ledger / mark     show what is left (with URLs) / mark an item done or blocked (quoted blocker)
  ingest            cards Claude read from a site (one line per card) -> pool.jsonl, ledger item done
  screen            gates + routing + pre-score over raw.json and pool.jsonl -> screened.json
  shortlist         dedupe vs both trackers, resolve aggregator finds to employer postings,
                    compact list of what Claude must open and score -> shortlist.json
  finalize          Claude's scores -> validated tracker rows for each tracker, merged files
  audit             ledger complete? (every item done, or blocked with a quoted blocker)
  report            the run report skeleton (funnel, coverage, rows added)

Maintenance: companies, registry-backfill, registry-check [network], h1b-build, h1b, stats,
calibrate, validate, state. Legacy v1 commands (fetch, run, pick, newgrad, dedupe, route) still work.

`python3 jobpipe.py <command> -h` for options.
"""
from __future__ import annotations

import argparse
import concurrent.futures as cf
import csv
import datetime as dt
import hashlib
import html
import io
import json
import math
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL_DIR = HERE.parent
CONFIG = json.loads(Path(os.environ.get("JOBPIPE_CONFIG", HERE / "config.json")).read_text())
_default_csv = SKILL_DIR / "references" / "companies.csv"
if not _default_csv.exists():  # standalone repo layout: companies.csv next to the script
    _default_csv = HERE / "companies.csv"
COMPANIES_CSV = Path(os.environ.get("JOBPIPE_COMPANIES", _default_csv))
UA = "jobpipe/2.0 (personal job search; contact via GitHub)"


def _now():
    v = os.environ.get("JOBPIPE_NOW")  # replay / tests: an ISO time
    if v:
        d = dt.datetime.fromisoformat(v.replace("Z", "+00:00"))
        return d if d.tzinfo else d.replace(tzinfo=dt.timezone.utc)
    return dt.datetime.now(dt.timezone.utc)


NOW = _now()
TODAY = NOW.astimezone(dt.timezone(dt.timedelta(hours=-4))).date()  # run day in Atlanta (EDT; good enough for labels)
RUBRIC = CONFIG["rubric"]
GATES = CONFIG["gates"]


# =========================================================================== utils
def log(*a):
    print(*a, file=sys.stderr)


def load_json(path, default=None):
    p = Path(path)
    if not p.exists():
        return [] if default is None else default
    txt = p.read_text(encoding="utf-8").strip()
    if not txt:
        return [] if default is None else default
    if p.suffix == ".jsonl":
        return [json.loads(x) for x in txt.splitlines() if x.strip()]
    data = json.loads(txt)
    if isinstance(data, dict) and default is None:  # accept {"roles":[...]} etc.
        for k in ("roles", "rows", "jobs", "candidates", "items", "data"):
            if isinstance(data.get(k), list):
                return data[k]
    return data


def load_rows(path):
    """Tracker exports may be JSON or CSV."""
    if str(path).lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))
    return load_json(path)


def dump_json(obj, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(obj, indent=1, ensure_ascii=False, default=str), encoding="utf-8")


def append_jsonl(rows, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "a", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False, default=str) + "\n")


def http(url, data=None, headers=None, timeout=25, retries=2):
    """(status, body, content-type). 429 and auth/404 errors return at once (never hammer a limit)."""
    hdrs = {"User-Agent": UA, "Accept": "application/json, text/html;q=0.9"}
    hdrs.update(headers or {})
    body = None
    if data is not None:
        body = json.dumps(data).encode()
        hdrs["Content-Type"] = "application/json"
    last = None
    for attempt in range(retries + 1):
        try:
            req = urllib.request.Request(url, data=body, headers=hdrs, method="POST" if body else "GET")
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.status, r.read().decode("utf-8", "replace"), r.headers.get("Content-Type", "")
        except urllib.error.HTTPError as e:
            if e.code in (429, 404, 410, 400, 401, 403, 405):
                return e.code, "", ""
            last = e
        except Exception as e:  # timeouts, resets, proxy refusals
            last = e
        time.sleep(1.5 * (attempt + 1))
    return 0, str(last), ""


def strip_html(s):
    s = html.unescape(s or "")
    s = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", s, flags=re.S | re.I)
    s = re.sub(r"<br\s*/?>|</p>|</li>|</div>|</h\d>", "\n", s, flags=re.I)
    s = re.sub(r"<[^>]+>", " ", s)
    s = html.unescape(s)
    s = re.sub(r"[ \t\r\f\v]+", " ", s)
    return re.sub(r"\n\s*\n+", "\n", s).strip()


def iso(ts):
    """epoch seconds/ms, ISO string, or 'October 27, 2026' -> ISO UTC string (or '')."""
    if ts in (None, ""):
        return ""
    try:
        if isinstance(ts, (int, float)) or str(ts).strip().isdigit():
            v = float(ts)
            if v > 1e12:
                v /= 1000
            return dt.datetime.fromtimestamp(v, dt.timezone.utc).isoformat()
        s = str(ts).strip().replace("Z", "+00:00")
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
            s += "T16:00:00+00:00"  # a bare date: midday US time, so it stays on the same calendar day
        try:
            d = dt.datetime.fromisoformat(s)
        except ValueError:
            d = None
            for fmt in ("%B %d, %Y", "%b %d, %Y", "%m/%d/%Y", "%Y-%m-%d", "%d %B %Y"):
                try:
                    d = dt.datetime.strptime(s, fmt).replace(hour=12)
                    break
                except ValueError:
                    continue
            if d is None:
                return ""
        if d.tzinfo is None:
            d = d.replace(tzinfo=dt.timezone.utc)
        return d.astimezone(dt.timezone.utc).isoformat()
    except Exception:
        return ""


def age_hours(iso_s):
    if not iso_s:
        return None
    try:
        t = dt.datetime.fromisoformat(str(iso_s).strip().replace("Z", "+00:00"))
        if t.tzinfo is None:
            t = t.replace(tzinfo=dt.timezone.utc)  # runstate keeps minute-precision UTC without an offset
        return (NOW - t).total_seconds() / 3600
    except Exception:
        return None


_UNIT_H = {"m": 1 / 60, "min": 1 / 60, "mins": 1 / 60, "minute": 1 / 60, "minutes": 1 / 60, "h": 1, "hr": 1, "hrs": 1,
           "hour": 1, "hours": 1, "d": 24, "day": 24, "days": 24, "w": 168, "wk": 168, "wks": 168, "week": 168,
           "weeks": 168, "mo": 720, "month": 720, "months": 720}
AGE_RX = re.compile(r"(?<![\w$])(\d{1,3})\s*(\+)?\s*(minutes?|mins?|m|hours?|hrs?|h|days?|d|weeks?|wks?|w|months?|mo)\b(?:\s*ago)?", re.I)


def parse_age(text):
    """Relative posting age -> hours (SKILL.md rule 19). 'Today'/'Just posted' 0, 'Yesterday' 24,
    '3h', '2 days ago', 'Posted 30+ Days Ago' (=> > 30 days), '1w'. None when no age is present."""
    t = (text or "").strip().lower()
    if not t:
        return None
    if re.search(r"\bjust (?:now|posted)\b|\bmoments? ago\b|\bposted today\b|^today$|\btoday\b", t):
        return 0.0
    if "yesterday" in t:
        return 24.0
    m = AGE_RX.search(t)
    if not m:
        return None
    n, plus, unit = int(m.group(1)), m.group(2), m.group(3).lower()
    h = n * _UNIT_H.get(unit, _UNIT_H.get(unit.rstrip("s"), 24))
    return h + (24 if plus else 0)


def workday_posted_on(text):
    """'Posted Today' / 'Posted Yesterday' / 'Posted 3 Days Ago' / 'Posted 30+ Days Ago' -> ISO."""
    a = parse_age(text)
    if a is None:
        return ""
    d = NOW - dt.timedelta(hours=a)
    return d.replace(minute=0, second=0, microsecond=0).isoformat()


def rx_any(patterns, text):
    for p in patterns:
        if re.search(p, text, re.I):
            return p
    return None


def wb(words):
    return [r"(?<![a-z])" + re.escape(w.strip()) + r"(?![a-z])" for w in words]


def short(s, n):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[: n - 1] + "…"


# =========================================================================== normalization (duplicate rule)
COMPANY_DROP = re.compile(r"\((?:posted via[^)]*|[^)]*)\)|,?\s*\b(inc|llc|ltd|corp|corporation|co|plc|l\.p\.|lp|n\.a\.|holdings|group|technologies|technology)\b\.?", re.I)
PARENT_ALIASES = {
    "lexisnexis risk solutions": "relx", "lexisnexis": "relx", "relx": "relx",
    "capital one services": "capital one", "jpmorganchase": "jpmorgan chase", "jpmorgan": "jpmorgan chase",
    "jp morgan chase": "jpmorgan chase", "j p morgan": "jpmorgan chase", "jpmorgan chase & co": "jpmorgan chase",
    "amazon web services": "amazon", "aws": "amazon", "amazon / aws": "amazon", "amazon com services": "amazon",
    "intuit mailchimp": "intuit", "mailchimp": "intuit", "intuit credit karma": "intuit", "credit karma": "intuit",
    "block, inc": "block", "square": "block", "cash app": "block", "rocket": "rocket companies",
    "mckinsey & company": "mckinsey", "mckinsey (quantumblack)": "mckinsey", "quantumblack": "mckinsey",
    "boston consulting group": "bcg", "bcg (bcg x)": "bcg", "bcg x": "bcg", "bain & company": "bain",
    "ernst & young": "ey", "pricewaterhousecoopers": "pwc", "deloitte consulting": "deloitte",
    "intercontinental exchange": "ice", "intercontinental exchange (ice)": "ice",
    "cox enterprises / cox automotive": "cox", "cox automotive": "cox", "cox enterprises": "cox",
    "the home depot": "home depot", "the coca-cola company": "coca-cola", "coca cola": "coca-cola", "meta platforms": "meta",
    "bank of america merrill lynch": "bank of america", "bofa": "bank of america", "wells fargo bank": "wells fargo",
    "citigroup": "citi", "citibank": "citi", "goldman sachs & co": "goldman sachs", "morgan stanley services": "morgan stanley",
    "fidelity": "fidelity investments", "fmr": "fidelity investments", "s&p global": "s&p global", "spgi": "s&p global",
    "sig": "susquehanna", "susquehanna international group": "susquehanna", "d e shaw": "d. e. shaw", "de shaw": "d. e. shaw",
    "imc": "imc trading", "imc financial markets": "imc trading", "hrt": "hudson river trading", "two sigma investments": "two sigma",
    "citadel securities": "citadel securities", "google llc": "google", "alphabet": "google", "microsoft corporation": "microsoft",
    "nvidia corporation": "nvidia", "walmart": "walmart & sam's club", "sam's club": "walmart & sam's club",
}


def norm_company(s):
    s0 = re.sub(r"\s+", " ", (s or "").strip().lower())
    if s0 in PARENT_ALIASES:
        return PARENT_ALIASES[s0]
    s = COMPANY_DROP.sub(" ", s0)
    s = re.sub(r"[^a-z0-9& ]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"^the ", "", s)
    return PARENT_ALIASES.get(s, s)


def norm_title(s):
    s = (s or "").lower()
    s = re.split(r"\s+[–—-]\s+(?:hybrid|remote|onsite|on-site)\b", s)[0]
    s = re.sub(r"\((?:hybrid|remote|onsite|on-site)[^)]*\)", " ", s)
    s = s.replace("sr.", "senior").replace("sr ", "senior ")
    s = re.sub(r"\b(i{1,3}|iv)\b", lambda m: str(len(m.group(1))) if m.group(1) != "iv" else "4", s)
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_city(s):
    s = (s or "").lower()
    if "remote" in s:
        return "remote"
    s = s.split(";")[0].split("|")[0].split(",")[0]
    s = re.sub(r"[^a-z ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_req(s):
    s = re.sub(r"[^a-z0-9]", "", str(s or "").lower())
    return s if len(s) >= 3 else ""


def fingerprint(company, title, location):
    key = f"{norm_company(company)}|{norm_title(title)}|{norm_city(location)}"
    return hashlib.sha1(key.encode()).hexdigest()[:12]


def title_tokens(t):
    stop = {"the", "a", "an", "and", "of", "for", "in", "to", "with", "at", "us", "remote", "hybrid", "onsite"}
    return {w for w in norm_title(t).split() if w not in stop and len(w) > 1}


def title_sim(a, b):
    x, y = title_tokens(a), title_tokens(b)
    return len(x & y) / max(1, len(x | y))


# =========================================================================== companies
CSV_FIELDS = ["company", "sector", "startup", "atlanta", "careers_url", "ats", "ats_slug", "workday_url", "search_url",
              "aliases", "status", "last_checked", "notes", "gt_fair", "fair_opt", "fair_sponsor", "source", "added",
              "lca_target", "lca_fy"]


def read_companies_csv(path=None):
    with open(path or COMPANIES_CSV, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        for k in CSV_FIELDS:
            r.setdefault(k, "")
    return rows


def write_companies_csv(rows, path=None):
    with open(path or COMPANIES_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def company_sectors(r):
    s = {r.get("sector", "")}
    if r.get("startup") == "yes":
        s.add("startups")
    return s


def load_companies(sectors=None, fair_only=False):
    rows = read_companies_csv()
    if sectors:
        want = set(sectors)
        rows = [r for r in rows if company_sectors(r) & want]
    if fair_only:
        rows = [r for r in rows if r.get("gt_fair")]
    rank = lambda r: (0 if r.get("fair_sponsor") == "yes" else 1 if r.get("gt_fair") else 2,
                      0 if r.get("atlanta") == "yes" else 1)
    return sorted(rows, key=rank)


_COMPANY_INDEX = None


def company_index():
    """norm name (and aliases) -> companies.csv row."""
    global _COMPANY_INDEX
    if _COMPANY_INDEX is None:
        idx = {}
        for r in read_companies_csv():
            idx[norm_company(r["company"])] = r
            for a in (r.get("aliases") or "").split(";"):
                if a.strip():
                    idx[norm_company(a)] = r
        _COMPANY_INDEX = idx
    return _COMPANY_INDEX


def company_row(name):
    return company_index().get(norm_company(name))


def has_feed(c):
    ats = (c.get("ats") or "").lower()
    if c.get("status") in ("no_feed", "broken"):
        return False
    if ats == "workday":
        return bool(c.get("workday_url"))
    if ats == "amazon":
        return True
    return ats in FETCHERS and bool(c.get("ats_slug"))


# =========================================================================== fetchers (employer boards)
def job(company_row_, **kw):
    c = company_row_ or {}
    d = dict(company=c.get("company", ""), sector=c.get("sector", ""), startup=c.get("startup", ""),
             atlanta=c.get("atlanta", ""), gt_fair=c.get("gt_fair", ""), fair_sponsor=c.get("fair_sponsor", ""),
             source="Employer job board", source_group="Employer job board", title="", location="", url="",
             apply_url="", req_id="", posted_at="", updated_at="", age_text="", ats="", description="",
             employment_type="", pay="")
    d.update(kw)
    d["description"] = (d.get("description") or "")[:8000]
    if not d.get("apply_url") and d.get("source_group") in ("Employer job board", "Company website"):
        d["apply_url"] = d.get("url", "")
    return d


def fetch_greenhouse(c, eu=False):
    host = "boards-api.eu.greenhouse.io" if eu or c.get("ats") == "greenhouse_eu" else "boards-api.greenhouse.io"
    st, body, _ = http(f"https://{host}/v1/boards/{c['ats_slug']}/jobs?content=true")
    if st != 200:
        return st, []
    out = []
    for j in json.loads(body).get("jobs", []):
        out.append(job(c, ats="greenhouse", title=j.get("title", ""), location=(j.get("location") or {}).get("name", ""),
                       url=j.get("absolute_url", ""), req_id=j.get("requisition_id") or str(j.get("id", "")),
                       posted_at=iso(j.get("first_published")), updated_at=iso(j.get("updated_at")),
                       description=strip_html(j.get("content", ""))))
    return st, out


def fetch_greenhouse_eu(c):
    return fetch_greenhouse(c, eu=True)


def fetch_lever(c):
    st, body, _ = http(f"https://api.lever.co/v0/postings/{c['ats_slug']}?mode=json")
    if st != 200:
        return st, []
    out = []
    for j in json.loads(body):
        cat = j.get("categories") or {}
        sal = j.get("salaryRange") or {}
        pay = f"{sal.get('min','')}-{sal.get('max','')} {sal.get('currency','')}".strip(" -") if sal else ""
        out.append(job(c, ats="lever", title=j.get("text", ""),
                       location=cat.get("location") or ", ".join(cat.get("allLocations") or []),
                       url=j.get("hostedUrl", ""), req_id=j.get("id", ""), posted_at=iso(j.get("createdAt")),
                       employment_type=cat.get("commitment", ""), pay=pay,
                       description=(j.get("descriptionPlain") or "") + " " + " ".join(
                           strip_html(x.get("content", "")) for x in j.get("lists") or [])
                       + " " + (j.get("additionalPlain") or "")))
    return st, out


def fetch_ashby(c):
    st, body, _ = http(f"https://api.ashbyhq.com/posting-api/job-board/{c['ats_slug']}?includeCompensation=true")
    if st != 200:
        return st, []
    out = []
    for j in json.loads(body).get("jobs", []):
        if j.get("isListed") is False:
            continue
        locs = [j.get("location") or ""] + [x.get("location", "") for x in j.get("secondaryLocations") or []]
        comp = (j.get("compensation") or {}).get("compensationTierSummary", "")
        out.append(job(c, ats="ashby", title=j.get("title", ""), location="; ".join(l for l in locs if l),
                       url=j.get("jobUrl", ""), req_id=j.get("id", ""), posted_at=iso(j.get("publishedAt")),
                       employment_type=j.get("employmentType", ""), pay=comp,
                       description=j.get("descriptionPlain") or strip_html(j.get("descriptionHtml", ""))))
    return st, out


def fetch_smartrecruiters(c):
    out, offset, st = [], 0, 0
    while offset < 1000:
        st, body, _ = http(f"https://api.smartrecruiters.com/v1/companies/{c['ats_slug']}/postings?limit=100&offset={offset}")
        if st != 200:
            break
        d = json.loads(body)
        for j in d.get("content", []):
            loc = j.get("location") or {}
            out.append(job(c, ats="smartrecruiters", title=j.get("name", ""),
                           location=", ".join(x for x in [loc.get("city"), loc.get("region"), loc.get("country")] if x),
                           url=f"https://jobs.smartrecruiters.com/{c['ats_slug']}/{j.get('id')}",
                           req_id=j.get("refNumber") or j.get("id", ""), posted_at=iso(j.get("releasedDate"))))
        if len(d.get("content", [])) < 100:
            break
        offset += 100
    return (200 if out else st), out


WD_RE = re.compile(r"https?://([a-z0-9-]+)\.(wd\d+)\.myworkdayjobs\.com/(?:[a-z]{2}-[A-Z]{2}/)?([A-Za-z0-9_-]+)")


def workday_queries():
    qs = []
    for prog in ("professional", "newgrad"):
        for fam in CONFIG["query_families"][prog].values():
            qs += fam["each"]
    return list(dict.fromkeys(qs))


def fetch_workday(c, queries=None, per_query=20):
    m = WD_RE.search(c.get("workday_url") or "")
    if not m:
        return 0, []
    tenant, wd, site = m.groups()
    api = f"https://{tenant}.{wd}.myworkdayjobs.com/wday/cxs/{tenant}/{site}/jobs"
    base = f"https://{tenant}.{wd}.myworkdayjobs.com/{site}"
    queries = queries or workday_queries()
    seen, out, st = set(), [], 0
    for q in queries:
        st, body, _ = http(api, data={"appliedFacets": {}, "limit": per_query, "offset": 0, "searchText": q})
        if st == 429:
            break
        if st != 200:
            continue
        for j in json.loads(body).get("jobPostings", []):
            path = j.get("externalPath", "")
            if not path or path in seen:
                continue
            seen.add(path)
            bullets = j.get("bulletFields") or []
            out.append(job(c, ats="workday", title=j.get("title", ""), location=j.get("locationsText", ""),
                           url=base + path, req_id=bullets[0] if bullets else "",
                           posted_at=workday_posted_on(j.get("postedOn", "")), age_text=j.get("postedOn", ""),
                           workday_api=f"https://{tenant}.{wd}.myworkdayjobs.com/wday/cxs/{tenant}/{site}{path}"))
        time.sleep(0.8)
    return (200 if out else st), out


def enrich_workday(j):
    """Second call for one Workday job: full description and exact start date (only for screened jobs)."""
    if not j.get("workday_api"):
        return j
    st, body, _ = http(j["workday_api"], retries=1)
    if st != 200:
        return j
    try:
        info = json.loads(body).get("jobPostingInfo", {})
    except Exception:
        return j
    j["description"] = strip_html(info.get("jobDescription", ""))[:8000]
    if info.get("startDate"):
        exact = iso(info["startDate"])
        if exact and (not j.get("posted_at") or exact < j["posted_at"]):  # stricter of the two (rule 19)
            j["posted_at"] = exact
    j["employment_type"] = info.get("timeType", "") or j.get("employment_type", "")
    j["req_id"] = info.get("jobReqId") or j.get("req_id", "")
    return j


def fetch_amazon(c, queries=None):
    out, seen, st = [], set(), 0
    for q in (queries or workday_queries()):
        url = ("https://www.amazon.jobs/en/search.json?" + urllib.parse.urlencode(
            {"base_query": q, "loc_query": "United States", "country": "USA", "result_limit": 50, "sort": "recent"}))
        st, body, _ = http(url)
        if st == 429:
            break
        if st != 200:
            continue
        for j in json.loads(body).get("jobs", []):
            jid = str(j.get("id_icims") or j.get("id") or "")
            if jid in seen:
                continue
            seen.add(jid)
            desc = " ".join(strip_html(j.get(k, "")) for k in ("description", "basic_qualifications", "preferred_qualifications"))
            out.append(job(c, ats="amazon", title=j.get("title", ""), location=j.get("normalized_location") or j.get("location", ""),
                           url="https://www.amazon.jobs" + (j.get("job_path") or ""), req_id=jid,
                           posted_at=iso(j.get("posted_date")), description=desc))
        time.sleep(0.8)
    return (200 if out else st), out



ORACLE_RE = re.compile(r"https?://([a-z0-9.-]+\.oraclecloud\.com)/hcmUI/CandidateExperience/[a-z_-]+/sites/([A-Za-z0-9_]+)", re.I)


def fetch_oracle(c, queries=None, per_query=25):
    """Oracle Recruiting Cloud (JPMorgan, Goldman Sachs, AmEx, ...). ats_slug = 'host|SITE'."""
    host, _, site = (c.get("ats_slug") or "").partition("|")
    if not host or not site:
        return 0, []
    out, seen, st = [], set(), 0
    for q in (queries or workday_queries()):
        finder = (f"findReqs;siteNumber={site},limit={per_query},keyword=\"{q}\",sortBy=POSTING_DATES_DESC")
        url = (f"https://{host}/hcmRestApi/resources/latest/recruitingCEJobRequisitions?onlyData=true"
               f"&expand=requisitionList.secondaryLocations&finder=" + urllib.parse.quote(finder, safe="=;,"))
        st, body, _ = http(url)
        if st == 429:
            break
        if st != 200:
            continue
        try:
            items = json.loads(body).get("items") or [{}]
        except Exception:
            continue
        for j in items[0].get("requisitionList") or []:
            jid = str(j.get("Id") or "")
            if not jid or jid in seen:
                continue
            seen.add(jid)
            locs = [j.get("PrimaryLocation") or ""] + [x.get("Name", "") for x in j.get("secondaryLocations") or []]
            out.append(job(c, ats="oracle", title=j.get("Title", ""), location="; ".join(l for l in locs if l),
                           url=f"https://{host}/hcmUI/CandidateExperience/en/sites/{site}/job/{jid}", req_id=jid,
                           posted_at=iso(j.get("PostedDate")), description=strip_html(j.get("ShortDescriptionStr") or ""),
                           oracle_detail=f"https://{host}/hcmRestApi/resources/latest/recruitingCEJobRequisitionDetails?expand=all"
                                         f"&onlyData=true&finder=" + urllib.parse.quote(f'ById;Id="{jid}",siteNumber={site}', safe="=;,")))
        time.sleep(0.8)
    return (200 if out else st), out


def enrich_oracle(j):
    if not j.get("oracle_detail"):
        return j
    st, body, _ = http(j["oracle_detail"], retries=1)
    if st != 200:
        return j
    try:
        it = (json.loads(body).get("items") or [{}])[0]
    except Exception:
        return j
    txt = " ".join(strip_html(it.get(k) or "") for k in ("ExternalDescriptionStr", "ExternalQualificationsStr", "ExternalResponsibilitiesStr"))
    if txt.strip():
        j["description"] = txt[:8000]
    return j

FETCHERS = {"oracle": fetch_oracle, "greenhouse": fetch_greenhouse, "greenhouse_eu": fetch_greenhouse_eu, "lever": fetch_lever,
            "ashby": fetch_ashby, "smartrecruiters": fetch_smartrecruiters, "workday": fetch_workday,
            "amazon": fetch_amazon}


def fetch_company(c):
    """One company, every board it has: ats_slug / workday_url may hold two values separated by ';'
    (e.g. a main board and a campus board)."""
    ats = (c.get("ats") or "").lower()
    if not has_feed(c):
        return {"company": c["company"], "status": "no_feed", "jobs": []}
    key = "workday_url" if ats == "workday" else "ats_slug"
    jobs, statuses = [], []
    for v in [x.strip() for x in (c.get(key) or "amazon").split(";") if x.strip()]:
        try:
            st, got = FETCHERS[ats]({**c, key: v})
        except Exception as e:  # a parser surprise on one board never stops the run
            st, got = f"error: {type(e).__name__}", []
        statuses.append(st)
        jobs += got
    st = 200 if 200 in statuses else statuses[0]
    seen, uniq = set(), []
    for j in jobs:
        k = j.get("url") or (j.get("title"), j.get("location"))
        if k not in seen:
            seen.add(k)
            uniq.append(j)
    return {"company": c["company"], "status": st, "jobs": uniq}


# =========================================================================== fetchers (lists: GitHub, HN)
SIMPLIFY_URL = "https://raw.githubusercontent.com/SimplifyJobs/New-Grad-Positions/dev/.github/scripts/listings.json"
SPEEDYAPPLY = ["https://raw.githubusercontent.com/speedyapply/2027-AI-College-Jobs/main/NEW_GRAD_USA.md",
               "https://raw.githubusercontent.com/speedyapply/2027-SWE-College-Jobs/main/NEW_GRAD_USA.md"]
HN_SEARCH = "https://hn.algolia.com/api/v1/search_by_date?tags=story,author_whoishiring&query=who%20is%20hiring"


def list_job(company, **kw):
    c = company_row(company) or {"company": company}
    d = job(c, **kw)
    d["company"] = company or c.get("company", "")
    return d


def fetch_simplify(path=None):
    if path:
        data = json.loads(Path(path).read_text())
    else:
        st, body, _ = http(SIMPLIFY_URL, timeout=60)
        if st != 200:
            return st, []
        data = json.loads(body)
    out = []
    for r in data:
        if not r.get("active") or r.get("is_visible") is False:
            continue
        sp = str(r.get("sponsorship") or "")
        desc = ("U.S. citizenship is required. " if "Citizenship" in sp else "") + \
               ("We do not sponsor visas. " if "Does Not Offer" in sp else "") + \
               ("PhD required. " if (r.get("degrees") and all("phd" in d.lower() for d in r["degrees"])) else "")
        out.append(list_job(r.get("company_name", ""), title=r.get("title", ""), location="; ".join(r.get("locations") or []),
                            url=r.get("url", ""), apply_url=r.get("url", ""), posted_at=iso(r.get("date_posted")),
                            req_id=str(r.get("id", ""))[:40], source="SimplifyJobs GitHub", source_group="Simplify Jobs",
                            program_hint="newgrad", list_category=r.get("category", ""), description=desc,
                            sponsorship_note=sp))
    return 200, out


SPEEDY_ROW = re.compile(r"^\|\s*(?P<company>.+?)\s*\|\s*(?P<title>.+?)\s*\|\s*(?P<loc>.+?)\s*\|(?P<rest>.*)\|\s*$")


def parse_speedyapply(md, source_url=""):
    out, section = [], ""
    for line in md.splitlines():
        if line.startswith("### "):
            section = line[4:].strip()
        m = SPEEDY_ROW.match(line)
        if not m or m.group("company").strip().lower() in ("company", "---"):
            continue
        cells = [x.strip() for x in m.group("rest").split("|")]
        href = re.search(r'href="(https?://[^"]+)"', m.group("rest"))
        if not href:
            continue
        age = cells[-1] if cells else ""
        pay = next((x for x in cells if "$" in x), "")
        company = strip_html(m.group("company"))
        out.append(list_job(company, title=strip_html(m.group("title")), location=strip_html(m.group("loc")),
                            url=href.group(1), apply_url=href.group(1), age_text=age,
                            posted_at=workday_posted_on(age), pay=pay, source="speedyapply GitHub",
                            source_group="Company website", program_hint="newgrad", list_category=section,
                            list_url=source_url))
    return out


def fetch_speedyapply(paths=None):
    out, st = [], 0
    for i, u in enumerate(SPEEDYAPPLY):
        if paths:
            md = Path(paths[i]).read_text() if i < len(paths) else ""
            st = 200
        else:
            st, md, _ = http(u, timeout=60)
            if st != 200:
                continue
        out += parse_speedyapply(md, u)
    return (200 if out else st), out


HN_TITLE_WORDS = ["data scien", "data analy", "analytics", "machine learning", "ml engineer", "ai engineer", "applied ai",
                  "applied scientist", "quant", "risk", "forward deployed", "solutions engineer", "solutions architect",
                  "founding engineer", "member of technical staff", "new grad", "software engineer", "llm", "genai"]


def fetch_hn(window_h=168, story_json=None, comments_json=None):
    """Top-level comments of the newest 'Ask HN: Who is hiring?' thread posted inside the window."""
    if story_json is not None:
        story = story_json
    else:
        st, body, _ = http(HN_SEARCH)
        if st != 200:
            return st, []
        hits = [h for h in json.loads(body).get("hits", []) if "who is hiring" in (h.get("title") or "").lower()]
        if not hits:
            return 404, []
        story = hits[0]
    sid = story.get("objectID") or story.get("story_id")
    cutoff = int((NOW - dt.timedelta(hours=window_h)).timestamp())
    comments, page = [], 0
    if comments_json is not None:
        comments = comments_json
    else:
        while page < 40:
            st, body, _ = http(f"https://hn.algolia.com/api/v1/search_by_date?tags=comment,story_{sid}"
                               f"&numericFilters=created_at_i>{cutoff}&hitsPerPage=100&page={page}")
            if st != 200:
                break
            d = json.loads(body)
            comments += d.get("hits", [])
            if page + 1 >= d.get("nbPages", 0):
                break
            page += 1
    out = []
    for c in comments:
        if str(c.get("parent_id")) != str(sid):  # top-level only
            continue
        text = strip_html(c.get("comment_text") or "")
        first = text.split("\n")[0]
        low = text.lower()
        if not any(w in low for w in HN_TITLE_WORDS):
            continue
        parts = [p.strip() for p in re.split(r"\s\|\s", first)]
        company = parts[0] if parts else ""
        role = next((p for p in parts[1:] if any(w in p.lower() for w in HN_TITLE_WORDS)), parts[1] if len(parts) > 1 else "")
        loc = next((p for p in parts[1:] if re.search(r"remote|onsite|hybrid|, [A-Z]{2}\b|new york|san francisco|nyc|sf\b|usa|us\b", p, re.I)), "")
        links = re.findall(r"https?://[^\s<>\"')]+", text)
        emp = next((u for u in links if any(h in u for h in JOB_HOSTS) or re.search(r"/(?:jobs?|careers?|positions?)\b", u)), "")
        out.append(list_job(company, title=role[:120], location=loc, url=f"https://news.ycombinator.com/item?id={c.get('objectID')}",
                            apply_url=emp, posted_at=iso(c.get("created_at_i") or c.get("created_at")),
                            source="Ask HN", source_group="Ask HN", description=text[:6000],
                            hn_links=links[:6], needs_employer_link=not bool(emp)))
    return 200, out


# =========================================================================== collect
def in_window(j, window_h):
    """False only for a posting whose date is known and older than its window; a posting with no date is kept
    (screen and finalize decide)."""
    age = age_hours(j.get("posted_at"))
    return age is None or age <= window_for_job(j, window_h)


def do_collect(window_h=72, workers=8, enrich=True, only=None, offline=None, prefilter=True):
    """Fetch every machine-readable source. `offline` = {'simplify': path, 'speedyapply': [paths]} for tests.
    With `prefilter`, raw.json keeps only postings inside the time window (every board lists ALL open roles, so
    the unfiltered file is tens of thousands of postings and the repo history would grow every day)."""
    only = set(only or ["feeds", "simplify", "speedyapply", "hn"])
    offline = offline or {}
    jobs, feeds, sources = [], [], {}
    if "feeds" in only:
        comps = load_companies()
        with cf.ThreadPoolExecutor(max_workers=workers) as ex:
            for r in ex.map(fetch_company, comps):
                feeds.append({"company": r["company"], "status": r["status"], "n": len(r["jobs"])})
                jobs.extend(r["jobs"])
        ok = sum(1 for f in feeds if f["status"] == 200)
        sources["Employer feeds"] = {"status": "ok" if ok else "failed", "companies_with_feed": sum(1 for f in feeds if f["status"] != "no_feed"),
                                     "feeds_ok": ok, "jobs": sum(f["n"] for f in feeds)}
        if enrich:  # Workday detail pages, only for roles that pass the title / seniority / window screen
            todo = [j for j in jobs if j.get("ats") in ("workday", "oracle") and quick_title_ok(j)
                    and (age_hours(j.get("posted_at")) or 0) <= max(window_h, 72)]
            with cf.ThreadPoolExecutor(max_workers=4) as ex:
                list(ex.map(lambda x: enrich_workday(x) if x.get("ats") == "workday" else enrich_oracle(x), todo[:800]))
            sources["Employer feeds"]["detail_pages_read"] = len(todo[:800])
    for key, fn, label in (("simplify", lambda: fetch_simplify(offline.get("simplify")), "SimplifyJobs GitHub"),
                           ("speedyapply", lambda: fetch_speedyapply(offline.get("speedyapply")), "speedyapply"),
                           ("hn", lambda: fetch_hn(CONFIG["window_hours"]["hn"]), "Ask HN")):
        if key in only:
            try:
                st, got = fn()
            except Exception as e:
                st, got = f"error: {type(e).__name__}: {e}", []
            sources[label] = {"status": "ok" if st == 200 else f"failed ({st})", "jobs": len(got)}
            jobs.extend(got)
    total_open = len(jobs)
    if prefilter:
        jobs = [j for j in jobs if in_window(j, window_h)]
    meta = {"generated_at": NOW.isoformat(), "window_hours": window_h, "sources": sources,
            "postings_seen": total_open, "postings_in_window": len(jobs),
            "feeds_failed": [f for f in feeds if f["status"] not in (200, "no_feed")],
            "no_feed": [f["company"] for f in feeds if f["status"] == "no_feed"]}
    return {"meta": meta, "feeds": feeds, "jobs": jobs}


def doctor():
    checks = {"job boards (Greenhouse API)": "https://boards-api.greenhouse.io/v1/boards/stripe/jobs",
              "job boards (Ashby API)": "https://api.ashbyhq.com/posting-api/job-board/ramp",
              "Workday": "https://wd1.myworkdayjobs.com/",
              "GitHub raw": SIMPLIFY_URL, "Hacker News API": HN_SEARCH}
    res = {}
    for k, u in checks.items():
        st, _, _ = http(u, retries=0, timeout=12)
        res[k] = st
    direct = res["job boards (Greenhouse API)"] == 200
    mode = "direct" if direct else ("github_only" if res["GitHub raw"] == 200 else "none")
    return mode, res


# =========================================================================== classification and gates
PROF = CONFIG["programs"]["professional"]
NG = CONFIG["programs"]["newgrad"]
TITLE_GROUP_RX = {g: wb(ws) for g, ws in PROF["title_groups"].items()}
NG_FAMILY_RX = {f: wb(ws) for f, ws in NG["families"].items()}
NG_MARKERS = NG["markers"]
SENIOR_RX = wb(GATES["seniority_title_exclude"])
SENIOR_ENG_RX = wb(GATES["senior_eng_exclude"])
SENIOR_OK_RX = wb(GATES["senior_ok_titles"])
NONUS_RX = wb(CONFIG["non_us_hints"])
USCITY_RX = wb(CONFIG["us_cities"]) + wb(CONFIG["us_location_hints"]) + wb(CONFIG["us_state_names"])
STATE_RX = re.compile(r"(?:,|\s|\()\s*(" + "|".join(CONFIG["us_states"]) + r")\b")
PAY_RX = re.compile(r"\$\s?(\d{2,3})(?:,(\d{3})|\.?\d?\s?[kK])\b")
NEWGRAD_SOURCES = {"SimplifyJobs GitHub", "speedyapply GitHub", "Simplify FAANG+", "newgrad-jobs.com"}
INTERN_WORDS = {"intern", "internship", "co-op", "coop", "apprentice", "summer analyst", "summer associate"}


def title_group(title):
    t = title.lower()
    for g, pats in TITLE_GROUP_RX.items():
        if rx_any(pats, t):
            return g
    return ""


def ng_family(title):
    t = title.lower()
    for f in ("mle", "ai", "quant", "fde", "sde"):
        if rx_any(NG_FAMILY_RX[f], t):
            return f
    return ""


def is_newgrad_title(title):
    t = title.lower()
    return bool(rx_any(NG_MARKERS, t)) and bool(ng_family(t))


def classify(j):
    """-> (program, family) or ('', '') when off-target."""
    t = j.get("title", "")
    marker = bool(rx_any(NG_MARKERS, t.lower())) and not re.search(r"\b(?:senior|sr\.?|staff|principal|lead)\b", t.lower())
    from_ng_list = j.get("program_hint") == "newgrad" or j.get("source") in NEWGRAD_SOURCES
    fam = ng_family(t)
    grp = title_group(t)
    if fam and (marker or from_ng_list) and grp not in ("data_science", "data_analytics", "consulting_strategy"):
        return "newgrad", fam
    if grp:
        return "professional", grp
    if fam and from_ng_list:
        return "newgrad", fam
    return "", ""


def quick_title_ok(j):
    prog, _ = classify(j)
    return bool(prog) and not seniority_reason(j.get("title", ""), prog, j.get("description", ""))


def seniority_reason(title, program="professional", text=""):
    """'' if the title passes the title gate; otherwise the reason."""
    t = " " + title.lower().replace("member of technical staff", "member of technical team") + " "
    for w in INTERN_WORDS:
        if re.search(r"(?<![a-z])" + re.escape(w) + r"(?![a-z])", t):
            return "intern/co-op"
    if program == "newgrad":
        hit = rx_any(GATES["newgrad_exclude_title"], t)
        if hit and re.fullmatch(r"ii|iii|2|3", re.search(hit, t, re.I).group(0).strip()) and re.search(r"\b(?:1|i)\s*(?:,|/|-|or|to)\s*(?:2|ii)\b", t):
            hit = None  # 'Level 1, 2, or 3' / 'Engineer I/II' include the entry level
        if hit and not re.search(r"\bnew grad\b.*\b(?:ii|2)\b", t):
            m_ = re.search(hit, t, re.I)
            return f"too senior for new grad (title: {(m_.group(0) if m_ else hit).strip()})"
    ok_senior = bool(rx_any(SENIOR_OK_RX, t))
    for w in GATES["seniority_title_exclude"]:
        if w in INTERN_WORDS:
            continue
        if re.search(r"(?<![a-z])" + re.escape(w) + r"(?![a-z])", t):
            if w == "manager" and manager_ic_ok(text):
                continue
            if w == "lead" and re.search(r"\blead(?:s|ing)? (?:generation|gen)\b", t):
                continue
            if w == "fellow" and re.search(GATES.get("early_program_fellow", "rotation"), t):
                continue  # 'AI Labs Rotational Fellow - Emerging Talent' is an early-career program
            return f"too senior (title: {w})"
    if not ok_senior and rx_any(SENIOR_ENG_RX, t) and not re.search(r"\b(?:engineer|scientist)\s+i\s*$", t.strip()):
        return "too senior (title: senior engineering)"
    if rx_any(GATES["level_exclude_patterns"], t):
        return "too senior (title: level III+ engineering)"
    return ""


def manager_ic_ok(text):
    """'Manager' title kept only when the posting is clearly an IC role asking <= 3 years (SKILL.md)."""
    if not text:
        return False
    low = text.lower()
    if rx_any(wb(CONFIG["profile"]["function_keywords"]["people_management"]), low):
        return False
    yrs = [m["n"] for m in year_mentions(text) if not m["preferred"]]
    return bool(yrs) and max(yrs) <= 3 and bool(re.search(r"individual contributor|\bic\b", low))


def is_us(location):
    l = (location or "").strip()
    if not l:
        return "unknown"
    low = l.lower()
    parts = re.split(r";|\||/| or ", low)
    us = bool(rx_any(USCITY_RX, low) or STATE_RX.search(l) or re.search(r"\bus\b|\busa\b|\bu\.s\.", low))
    nonus_all = all(rx_any(NONUS_RX, p) for p in parts if p.strip()) if parts else False
    if us:
        return "yes"
    if nonus_all or rx_any(NONUS_RX, low):
        return "no"
    return "unknown"


def stated_max_pay(text):
    vals = []
    for m in PAY_RX.finditer(text or ""):
        if m.group(2) is not None:
            v = int(m.group(1) + m.group(2))
        else:
            v = int(m.group(1)) * 1000
        if 20000 <= v <= 1000000:
            vals.append(v)
    if not vals:
        return None
    low = (text or "").lower()
    if re.search(r"/\s?h(?:ou)?r|per hour|hourly", low) and max(vals) < 300000:
        return None  # hourly rates: leave to Claude
    return max(vals)


# --------------------------------------------------------------------------- years by function
YEARS_RX = re.compile(r"(?P<n>\d{1,2})(?:\.\d)?\s*(?:\+|plus)?\s*(?:(?:-|–|to)\s*(?P<m>\d{1,2})\s*)?\+?\s*(?:years?|yrs?)\b"
                      r"(?P<after>(?:(?!\d{1,2}(?:\.\d)?\s*\+?\s*(?:years?|yrs?)\b)[^.;\n]){0,150})", re.I)
FUNC_RX = {f: wb(ws) for f, ws in CONFIG["profile"]["function_keywords"].items()}
FY = CONFIG["profile"]["function_years"]


def classify_function(ctx):
    hits = [f for f, pats in FUNC_RX.items() if rx_any(pats, ctx)]
    return hits


def year_mentions(text):
    out = []
    t = text or ""
    for m in YEARS_RX.finditer(t):
        n = int(m.group("n"))
        if n == 0 or n > 20:
            continue
        before = t[max(0, m.start() - 90): m.start()].lower()
        after = m.group("after").lower()
        ctx = before[-60:] + " " + after
        if re.search(r"years? (?:old|of age)|per year|a year\b|/year|years? ago|within \d+ years|year[- ]round|\bage\b", after[:25] + before[-15:]):
            continue
        if not re.search(r"experience|exp\b|work|working|professional|industry|background|of\s+\w+|in\s+\w+|building|developing|managing|leading|hands-on", ctx):
            continue
        look = t[max(0, m.start() - 500): m.start()].lower()
        pref_pos = max(look.rfind(k) for k in ("preferred", "nice to have", "bonus", "a plus", "desired", "ideally"))
        req_pos = max(look.rfind(k) for k in ("required", "requirements", "minimum", "basic qualifications", "must have", "you have", "qualifications", "what you"))
        preferred = pref_pos > req_pos or bool(re.search(r"\bpreferred\b|\ba plus\b|\bnice to have\b", after[:80]))
        deg = "master" if re.search(r"master|m\.s\.|\bms\b|graduate degree|advanced degree", before[-80:] + after[:40]) else \
              "phd" if re.search(r"ph\.?d|doctor", before[-60:] + after[:30]) else \
              "bachelor" if re.search(r"bachelor|b\.s\.|\bbs\b|undergraduate", before[-80:] + after[:40]) else ""
        # the noun right after "years of ..." names the function ("4+ years of data engineering experience
        # with SQL" is engineering, not analytics); the wider window is only a fallback
        primary = re.split(r"\b(?:with|using|including|such as|across|leveraging|utilizing)\b|[,;(]", after)[0][:70]
        funcs = classify_function(primary) or classify_function(after[:120] + " " + before[-40:])
        out.append({"n": n, "funcs": funcs, "preferred": preferred, "degree": deg, "text": short(t[m.start(): m.end()], 90)})
    return out


def years_analysis(text):
    """Requirement path best for her (Master's > none > Bachelor's; PhD paths ignored), every required
    mention in it counts (AND). Returns gate reason ('' = pass), experience points, evidence line, flags."""
    ms = [m for m in year_mentions(text) if not m["preferred"] and m["degree"] != "phd"]
    if not ms:
        return {"gate": "", "points": None, "ask": None, "ev": "", "flags": []}
    paths = defaultdict(list)
    for m in ms:
        paths[m["degree"]].append(m)
    chosen = paths.get("master") or paths.get("") or paths.get("bachelor")
    if paths.get("master") and paths.get(""):
        chosen = paths["master"] + paths[""]
    yg = GATES["years_gate"]
    worst_gap, gate, parts, has_parts, flags = -99.0, "", [], [], []
    for m in chosen:
        funcs = m["funcs"] or ["general"]
        best = max(funcs, key=lambda f: FY.get(f, 0))  # an "or" list: the row most in her favour
        if any(f in yg["hard_functions"] for f in funcs) and all(f in yg["hard_functions"] or FY.get(f, 0) < 1 for f in funcs):
            hard = max((f for f in funcs if f in yg["hard_functions"]), key=lambda f: FY.get(f, 0))
            gap = m["n"] - FY.get(hard, 0)
            if m["n"] >= 2 and gap > yg["hard_margin"] and not gate:
                gate = f"too senior: asks {m['n']} yrs {hard.replace('_', ' ')}, has {FY.get(hard, 0):g}"
            best = hard
        has = FY.get(best, 0)
        gap = (m["n"] - has) if m["n"] >= 2 else -1
        worst_gap = max(worst_gap, gap)
        parts.append(f"{m['n']} yrs {best.replace('_', ' ')}")
        has_parts.append(f"{has:g} ({best.replace('_', ' ')})")
    pts = 20
    for lim, p in yg["soft_points"]:
        if worst_gap <= lim:
            pts = p
            break
    if 1 < worst_gap <= 3:
        flags.append("Stretch: years gap")
    elif worst_gap > 3:
        flags.append(f"large years gap ({worst_gap:g})")
    ev = f"Asks: {', '.join(dict.fromkeys(parts))}. Has: {', '.join(dict.fromkeys(has_parts))}."
    return {"gate": gate, "points": pts, "ask": max(m["n"] for m in chosen), "ev": ev, "flags": flags}


# --------------------------------------------------------------------------- H-1B history
_H1B = None


def load_h1b(path=None):
    global _H1B
    if _H1B is None or path:
        p = Path(path or os.environ.get("JOBPIPE_H1B", HERE / "h1b.json"))
        _H1B = load_json(p, default={}) if p.exists() else {}
    return _H1B


def h1b_lookup(company):
    db = load_h1b()
    if not db:
        return None
    k = norm_company(company)
    if k in db:
        return db[k]
    cands = [x for x in db if x.startswith(k + " ") or k.startswith(x + " ")]
    if cands:
        return db[max(cands, key=len)]
    return None


def sponsorship_level(j):
    text = (j.get("description") or "") + " " + (j.get("sponsorship_note") or "")
    if rx_any(GATES["sponsor_yes_patterns"], text):
        return "explicit_yes", "posting offers sponsorship"
    sp = sponsor_lookup(j.get("company", ""))
    if sp:
        n = sum(sp["groups"].values())
        core = sp["groups"].get("data_science", 0) + sp["groups"].get("finance_risk", 0) + sp["groups"].get("analytics", 0)
        lvl = "h1b_100" if n >= 100 else "h1b_10" if n >= 10 else "h1b_1"
        return lvl, f"DOL LCA FY{sp['fy']}: {n} H-1B filings in her role codes ({core} DS/analytics/finance)"
    rec = h1b_lookup(j.get("company", ""))
    if rec:
        a = int(rec.get("approvals", 0))
        lvl = "h1b_100" if a >= 100 else "h1b_10" if a >= 10 else "h1b_1" if a >= 1 else "h1b_none"
        return lvl, f"USCIS FY{rec.get('fy','?')}: {a} H-1B approvals"
    if j.get("fair_sponsor") == "yes":
        return "fair_yes", "told GT career fair it sponsors"
    if j.get("h1b_history"):
        return ("h1b_none", j["h1b_history"]) if "no record" in j["h1b_history"].lower() else ("h1b_10", j["h1b_history"])
    return "unknown", "no sponsorship data (check MyVisaJobs)"


def parse_h1b_csv(text, fy_hint=""):
    """USCIS H-1B Employer Data Hub export (column names vary by year) -> {norm employer: {approvals, fy}}."""
    rd = csv.DictReader(io.StringIO(text.lstrip("﻿")), delimiter="\t" if text.count("\t") > text.count(",") else ",")
    cols = rd.fieldnames or []
    emp = next((c for c in cols if "employer" in c.lower() or "petitioner" in c.lower()), None)
    fy = next((c for c in cols if "fiscal" in c.lower() or c.lower().strip() == "fy"), None)
    appr = [c for c in cols if "approval" in c.lower()]
    out = {}
    if not emp or not appr:
        return out
    for r in rd:
        k = norm_company(r.get(emp, ""))
        if not k:
            continue
        n = 0
        for c in appr:
            try:
                n += int(str(r.get(c) or "0").replace(",", ""))
            except ValueError:
                pass
        year = (r.get(fy) or fy_hint or "").strip() if fy else fy_hint
        cur = out.get(k)
        if not cur or str(year) > str(cur["fy"]):
            out[k] = {"approvals": n, "fy": year}
        elif str(year) == str(cur["fy"]):
            cur["approvals"] += n
    return out


# --------------------------------------------------------------------------- screen one job
EXC_LABEL = {"no_sponsorship": "no sponsorship", "citizenship": "citizenship required",
             "clearance": "clearance required", "phd_required": "PhD required"}


def window_for_job(j, base):
    w = base
    if j.get("sector") == "proptech":
        w = max(w, CONFIG["window_hours"]["proptech"])
    if j.get("startup") == "yes":
        w = max(w, CONFIG["window_hours"]["startup"])
    if j.get("source") in ("SimplifyJobs GitHub", "speedyapply GitHub"):
        w = max(w, CONFIG["window_hours"]["github_lists"])
    if j.get("source") == "Ask HN":
        w = max(w, CONFIG["window_hours"]["hn"])
    return w


def screen_one(j, window_h, seen=None):
    j = dict(j)
    reasons, flags = list(j.get("exclude_reasons") or []), list(j.get("flags") or [])
    t = j.get("title", "")
    text = (j.get("description") or "") + " " + (j.get("snippet") or "")
    prog, fam = classify(j)
    m = get_matcher()
    jm = m.score(j.get("description") or "") if m else None
    j["jd_match"] = jm
    if not prog and jm and jm["match"] >= JDM.get("keep_other_title_at", 0.62) \
            and not rx_any(wb(JDM.get("other_title_negative", [])), " " + t.lower() + " "):
        prog, fam = "professional", "other"  # title outside every group, but the posting matches her resume
        flags.append(f"title outside target groups; kept on JD match {jm['match']:.2f}")
    j["program"], j["family"] = prog, fam
    j["tracker"] = "newgrad" if prog == "newgrad" else "main" if prog else ""
    j["title_group"] = fam
    if not prog:
        reasons.append("off-target title")
    sr = seniority_reason(t, prog or "professional", text)
    if sr:
        reasons.append(sr)
    us = is_us(j.get("location", ""))
    j["us"] = us
    if us == "no":
        reasons.append("non-US location")
    elif us == "unknown":
        flags.append("location unclear")
    # freshness: exact ATS date first, then a relative age (rule 19); stricter of the two when both exist
    a = age_hours(j.get("posted_at"))
    rel = parse_age(j.get("age_text", ""))
    if a is None and rel is not None:
        a = rel
        j["posted_at"] = (NOW - dt.timedelta(hours=rel)).isoformat()
        flags.append("relative age" + (f" from {j.get('source')}" if j.get("source") not in ("Employer job board",) else ""))
    elif a is not None and rel is not None:
        a = max(a, rel)
    j["age_hours"] = round(a, 1) if a is not None else None
    win = window_for_job(j, window_h)
    if a is None:
        flags.append("date unverified")
    elif a > win:
        reasons.append(f"stale ({a / 24:.1f} days)")
    if j.get("known_company"):
        reasons.append("known company: harvested directly")
    if j.get("reposted") or re.search(r"\breposted\b", " ".join(j.get("card_flags") or []), re.I):
        reasons.append("reposted")
    for kind, pats in GATES["exclude_patterns"].items():
        if rx_any(pats, text):
            reasons.append(EXC_LABEL[kind])
    low_c = (j.get("company") or "").lower()
    if any(re.search(r"(?<![a-z])" + re.escape(d) + r"(?![a-z])", low_c) for d in GATES["defense_employers"]):
        reasons.append("defense / clearance employer")
    if any(s in low_c for s in GATES["staffing_agencies"]):
        flags.append("staffing agency (client may be undisclosed; low priority)")
    if rx_any(GATES["phd_preferred_patterns"], text):
        flags.append("PhD preferred")
    if prog == "newgrad" and re.search(r"class of 2027|2027 grad|graduat\w* (?:in|between)[^.]{0,40}2027|expected graduation[^.]{0,30}2027", text.lower() + " " + t.lower()):
        flags.append("asks 2027 graduation (she graduated May 2026)")
    ya = years_analysis(text) if prog == "professional" else {"gate": "", "points": None, "ev": "", "flags": [], "ask": None}
    if ya["gate"]:
        reasons.append(ya["gate"])
    flags += ya["flags"]
    j["years_required"] = ya["ask"]
    j["ev_experience_auto"] = ya["ev"]
    j["_exp_pts"] = ya["points"]
    pay = stated_max_pay((j.get("pay") or "") + " " + text[:5000])
    j["pay_max_stated"] = pay
    if pay and pay < CONFIG["salary_floor"]:
        reasons.append(f"pay below floor (${pay:,})")
    if re.search(r"\bcontract\b|\bcontractor\b|\btemporary\b|\bc2h\b", ((j.get("employment_type") or "") + " " + t).lower()):
        flags.append("contract")
    if not (j.get("description") or "").strip():
        flags.append("no description yet (open the posting)")
    elif "sponsor" not in text.lower() and "visa" not in text.lower():
        flags.append("sponsorship not mentioned")
    j["fp"] = j.get("fp") or fingerprint(j.get("company"), t, j.get("location"))
    if seen is not None and not reasons:
        first = seen.get(j["fp"])
        if not first:
            seen[j["fp"]] = NOW.isoformat()
        elif (age_hours(first) or 0) > CONFIG["repost_after_hours"] and (j["age_hours"] or 0) < 48:
            reasons.append(f"reposted (first seen {first[:10]})")
        j["first_seen"] = seen.get(j["fp"])
    j["exclude_reasons"] = list(dict.fromkeys(reasons))
    j["flags"] = list(dict.fromkeys(flags))
    j["screen"] = "pass" if not j["exclude_reasons"] else "excluded"
    if j["screen"] == "pass":
        prescore(j)
    return j


# --------------------------------------------------------------------------- pre-score
OTHER_TECH = ["java", "golang", " go ", "rust", "kotlin", "swift", "c#", ".net", "ruby", "php", "scala", "c++", "react",
              "angular", "vue", "node.js", "typescript", "kubernetes", "terraform", "salesforce", "sap", "cobol",
              "verilog", "embedded c", "unity", "unreal"]


def skills_points(text, max_pts=30):
    low = " " + (text or "").lower() + " "
    if len(low.strip()) < 200:
        return round(max_pts * 0.6), "skills not checked (no description)"
    strong = [s for s in CONFIG["profile"]["skills"]["strong"] if s in low]
    some = [s.strip() for s in CONFIG["profile"]["skills"]["some"] if s in low]
    missing = [s.strip() for s in OTHER_TECH if s in low and s.strip() not in some]
    denom = len(strong) + 0.5 * len(some) + len(missing) + 1.5
    frac = (len(strong) + 0.5 * len(some)) / denom
    pts = round(max_pts * min(1.0, 0.3 + 0.8 * frac))
    ev = f"matches {', '.join(strong[:6]) or 'none of her core skills'}"
    if missing:
        ev += f"; not on resume: {', '.join(missing[:4])}"
    return pts, ev


def prescore(j):
    prog = j.get("program") or "professional"
    w = RUBRIC[prog]
    text = j.get("description") or ""
    low = (j.get("title", "") + " " + text).lower()
    comp, ev = {}, {}
    # role fit
    if prog == "professional":
        rf = RUBRIC["role_fit_professional"].get(j.get("family"), 15)
    else:
        rf = RUBRIC["role_fit_newgrad"].get(j.get("family"), 15)
        if j.get("family") == "sde":
            tl = j.get("title", "").lower()
            if rx_any(wb(NG["low_fit_words"]), tl):
                rf = 6
            elif rx_any(wb(NG["high_fit_words"]), tl):
                rf = 22
    comp["role_fit"] = min(w["role_fit"], rf)
    ev["role_fit"] = f"{j.get('family','?').replace('_', ' ')} track"
    # skills: resume <-> JD match when the posting text is there, keyword fallback otherwise
    jm = j.get("jd_match")
    if jm:
        comp["skills"] = round(w["skills"] * min(1.0, JDM["floor"] + JDM["span"] * jm["match"]))
        ev["skills"] = (f"has {len(jm['has'])}/{len(jm['asks'])} asked ({', '.join(jm['has'][:5]) or '-'})"
                        + (f"; missing {', '.join(jm['missing'][:4])}" if jm["missing"] else ""))
        j["resume_suggested"] = jm["resume"]
    else:
        comp["skills"], ev["skills"] = skills_points(text, w["skills"])
    # experience (professional) / eligibility (new grad)
    if prog == "professional":
        pts = j.get("_exp_pts")
        comp["experience"] = w["experience"] if pts is None else min(w["experience"], pts)
        if pts is None:
            comp["experience"] = round(w["experience"] * 0.85)
        ev["experience"] = j.get("ev_experience_auto") or "Asks: years not stated (check posting). Has: see table."
    else:
        e = w["eligibility"]
        if any("2027 graduation" in f for f in j.get("flags", [])):
            e -= 6
        if "PhD preferred" in j.get("flags", []):
            e -= 2
        comp["eligibility"] = max(0, e)
        ev["eligibility"] = "new-grad program; M.S. May 2026" + ("; asks 2027 grads" if e < w["eligibility"] - 3 else "")
    # domain & qualifications
    d = RUBRIC["domain_by_sector"].get(j.get("sector", ""), RUBRIC["domain_by_sector"][""])
    if any(x in low for x in RUBRIC["domain_words"]):
        d += 2
    if prog == "professional" and re.search(r"\bcfa\b|\bacca\b|\bfrm\b", low):
        d += 2
    comp["domain_quals"] = min(w["domain_quals"], d)
    ev["domain_quals"] = f"sector {j.get('sector') or 'unknown'}"
    # preferences = sponsorship + logistics
    split = RUBRIC["preferences_split"][prog]
    lvl, why = sponsorship_level(j)
    j["sponsorship_level"], j["h1b_history"] = lvl, (j.get("h1b_history") or why)
    sp = round(split["sponsorship"] * RUBRIC["sponsorship_levels"][lvl])
    lg = (2 if j.get("us") == "yes" else 1) + (1 if re.search(r"atlanta|remote", (j.get("location") or "").lower()) else 0)
    lg += 1 if (j.get("pay_max_stated") is None or j["pay_max_stated"] >= CONFIG["salary_floor"]) else 0
    lg += 1 if (j.get("age_hours") is not None and j["age_hours"] <= 24) else 0
    comp["preferences"] = min(w["preferences"], sp + min(split["logistics"], lg))
    ev["preferences"] = f"{why}; {j.get('location') or 'location ?'}"
    exp_key = "experience" if prog == "professional" else "eligibility"
    j["pre_components"] = {k: comp[k] for k in ("skills", exp_key, "role_fit", "domain_quals", "preferences")}
    j["pre_evidence"] = ev
    j["pre_score"] = sum(j["pre_components"].values())
    return j


def build_matcher(jobs, resumes=None):
    resumes = load_resumes() if resumes is None else resumes
    if not resumes:
        return None
    texts = [j.get("description") or "" for j in jobs if len(j.get("description") or "") >= 200]
    m = JDMatcher(resumes, texts)
    m.calibrate(texts)
    return m


def do_screen(jobs, window_h, seen=None, matcher="auto"):
    if matcher == "auto":
        matcher = build_matcher(jobs)
    set_matcher(matcher)
    return [screen_one(j, window_h, seen) for j in jobs]



# =========================================================================== resume <-> job description matching
JDM = CONFIG.get("jd_match", {})
STOP = set("""a an and are as at be been but by can for from has have if in into is it its may more most must not of on or our
out over per so such than that the their them then there these they this those to under up use used using via was we were what when
where which while who will with within without you your yours about across all also any able both each etc i e g including like
new one other own role team teams work working years year experience strong ability skills knowledge preferred required plus
help join company opportunity looking us based position job candidate candidates""".split())
WORD_RX = re.compile(r"[a-z][a-z0-9+#./-]*[a-z0-9+#]|[a-z]")


def _compile_lexicon():
    out = []
    for cls, terms in (JDM.get("lexicon") or {}).items():
        for canon, pats in terms.items():
            rx = [re.compile(p_, 0 if re.search(r"[A-Z]", p_.replace("\\b", "")) else re.I) for p_ in pats]
            out.append((cls, canon, rx))
    return out


LEXICON = _compile_lexicon()


def lexicon_terms(text):
    t = text or ""
    return {canon: cls for cls, canon, rxs in LEXICON if any(r.search(t) for r in rxs)}


def tokens(text):
    ws = [w.strip("./-") for w in WORD_RX.findall((text or "").lower())]
    ws = [w for w in ws if len(w) > 1 and w not in STOP and not w.isdigit()]
    return ws + [a + "_" + b for a, b in zip(ws, ws[1:])]


def load_resumes(folder=None):
    folder = Path(folder or os.environ.get("JOBPIPE_RESUMES", SKILL_DIR / "resumes"))
    if not folder.exists():
        return {}
    return {p_.name: p_.read_text(encoding="utf-8") for p_ in sorted(folder.glob("*.md"))}


class JDMatcher:
    """Scores a posting against each resume: requirement coverage (lexicon terms the posting asks for that
    the resume shows; terms in requirement lines count double) blended with TF-IDF similarity, ranked
    against the day's postings. What she has comes from the resumes, never from config."""

    def __init__(self, resumes, corpus=(), stats=None):
        self.resumes = resumes
        self.has = {n: set(lexicon_terms(t)) for n, t in resumes.items()}
        self.has_any = set().union(*self.has.values()) if self.has else set()
        if stats:
            self.df, self.n, self.cos_sorted = Counter(stats["df"]), stats["n"], stats["cos_sorted"]
        else:
            self.df, self.n = Counter(), 0
            for txt in list(corpus) + list(resumes.values()):
                self.df.update(set(tokens(txt)))
                self.n += 1
            self.cos_sorted = []
        self.rvec = {n: self._vec(t) for n, t in resumes.items()}

    def _vec(self, text):
        tf = Counter(tokens(text))
        v = {w: (1 + math.log(c)) * math.log((self.n + 1) / (self.df.get(w, 0) + 1)) for w, c in tf.items()}
        norm = math.sqrt(sum(x * x for x in v.values())) or 1.0
        return {w: x / norm for w, x in v.items()}

    def cosine(self, text):
        v = self._vec(text)
        return {n: sum(x * rv.get(w, 0.0) for w, x in v.items()) for n, rv in self.rvec.items()}

    def calibrate(self, texts):
        self.cos_sorted = sorted(max(self.cosine(t).values() or [0]) for t in texts if len(t or "") >= 200)

    def pct(self, cos):
        cs = self.cos_sorted
        if not cs:
            return min(1.0, cos / 0.3)
        import bisect
        return bisect.bisect_left(cs, cos) / len(cs)

    def stats(self):
        return {"df": {w: c for w, c in self.df.items() if c > 1}, "n": self.n, "cos_sorted": self.cos_sorted[::max(1, len(self.cos_sorted) // 400)]}

    def score(self, text):
        if not self.resumes or len(text or "") < 200:
            return None
        asks = lexicon_terms(text)
        req = lexicon_terms(" ".join(l for l in re.split(r"\n|(?<=[.;])\s+", text) if REQ_LINE.search(l)))
        cw = JDM.get("class_weight", {})
        best = None
        cos = self.cosine(text)
        for name, has in self.has.items():
            tot = got = 0.0
            for canon, cls in asks.items():
                w_ = cw.get(cls, 1.0) * (2.0 if canon in req else 1.0)
                tot += w_
                got += w_ if canon in has else 0.0
            n_asks = len(asks)
            cov = got / tot if tot else 0.5
            if n_asks < JDM.get("min_asks", 3):  # too few asks to judge: pull toward neutral
                cov = 0.5 + (cov - 0.5) * n_asks / JDM.get("min_asks", 3)
            pc = self.pct(cos[name])
            m = JDM["weights"]["coverage"] * cov + JDM["weights"]["similarity"] * pc
            if not best or m > best["match"]:
                best = {"match": round(m, 3), "coverage": round(cov, 3), "similarity_pct": round(pc, 3), "resume": name,
                        "asks": sorted(asks), "has": sorted(a for a in asks if a in has),
                        "missing": sorted((a for a in asks if a not in has), key=lambda a: (a not in req, a))}
        return best


_MATCHER = None


def get_matcher():
    return _MATCHER


def set_matcher(m):
    global _MATCHER
    _MATCHER = m

# =========================================================================== links, LinkedIn ID age, dedupe
AGGREGATOR_HOSTS = ("linkedin.com", "jobright.ai", "hiring.cafe", "hiringcafe.com", "builtin.com", "indeed.com",
                    "joinhandshake.com", "simplify.jobs", "trueup.io", "levels.fyi", "welcometothejungle.com",
                    "wellfound.com", "glassdoor.com", "ziprecruiter.com", "efinancialcareers.com", "symplicity.com",
                    "interstride.com", "airtable.com", "newgrad-jobs.com", "ycombinator.com/jobs", "workatastartup.com",
                    "news.ycombinator.com", "otta.com", "dice.com", "monster.com", "careerbuilder.com")
JOB_HOSTS = ("greenhouse.io", "lever.co", "ashbyhq.com", "myworkdayjobs.com", "myworkdaysite.com", "smartrecruiters.com",
             "oraclecloud.com", "icims.com", "taleo.net", "successfactors.com", "workable.com", "jobvite.com",
             "bamboohr.com", "eightfold.ai", "phenompeople.com", "avature.net", "recruitee.com", "rippling.com",
             "paylocity.com", "ultipro.com", "adp.com", "amazon.jobs", "gem.com", "pinpointhq.com", "breezy.hr",
             "teamtailor.com", "personio", "dover.com", "jobs.apple.com", "careers.google.com", "google.com/about/careers",
             "jobs.careers.microsoft.com", "metacareers.com")


def apply_link_problem(url):
    """'' if url looks like one specific employer posting; otherwise why not (SKILL.md rule 18)."""
    u = (url or "").strip()
    if not u.startswith("http"):
        return "no link"
    p = urllib.parse.urlparse(u)
    host, path = p.netloc.lower(), p.path.rstrip("/")
    full = host + path
    if any(h in full for h in AGGREGATOR_HOSTS):
        return f"aggregator page ({host}), not the employer posting"
    has_id = bool(re.search(r"\d{4,}|[A-Z]{1,4}[-_]?\d{3,}|[0-9a-f]{8}-[0-9a-f]{4}", u)) or "gh_jid=" in u or "jobid=" in u.lower()
    if any(h in host for h in JOB_HOSTS):
        return "" if (has_id or path.count("/") >= 2) else "ATS board home, not one posting"
    if path.count("/") <= 1 and not has_id:
        return "careers home page, not one posting"
    if not has_id and not re.search(r"/(job|jobs|position|positions|opening|openings|req|requisition|careers?|details|posting)/[^/]+", path, re.I):
        return "no job ID in link"
    return ""


LI_ID_RX = re.compile(r"linkedin\.com/jobs/(?:view|search)/?(?:[^/?#]*-)?(?:\?currentJobId=)?(\d{8,})")
LI_IDS_PER_DAY = 500_000  # Oct 2026: 4482.7M on Oct 24, 4483.8M on Oct 26


def linkedin_age_check(pool, max_days=3, ref_id=None):
    """Old LinkedIn job ID + fresh label = repost / bump (US Bank ~12 days, FHLB Indianapolis ~100 days)."""
    ids = []
    for j in pool:
        m = LI_ID_RX.search(j.get("url") or j.get("apply_link") or "")
        if m:
            j["linkedin_id"] = int(m.group(1))
            ids.append(j["linkedin_id"])
    top = ref_id or (max(ids) if ids else 0)
    for j in pool:
        lid = j.get("linkedin_id")
        if lid and top:
            est = (top - lid) / LI_IDS_PER_DAY
            j["linkedin_est_age_days"] = round(est, 1)
            if est > max_days:
                j.setdefault("exclude_reasons", []).append(f"reposted / old LinkedIn ID (~{est:.0f} days older than newest)")
                j["screen"] = "excluded"
    return pool


def tracker_index(rows):
    reqs, fps, by_company = {}, {}, defaultdict(list)
    for r in rows:
        title = r.get("title") or r.get("job_title") or ""
        rid = norm_req(r.get("req_id") or r.get("requisition_id") or "")
        status = r.get("status", "")
        nc = norm_company(r.get("company"))
        if rid:
            reqs[(nc, rid)] = status
        fresh = fingerprint(r.get("company"), title, r.get("location"))
        for fp in {r.get("fp") or fresh, fresh}:  # stored fp may come from an older normalizer
            fps[fp] = status
        fps.setdefault(fingerprint(r.get("company"), title, ""), status)
        for link in (r.get("apply_link"), r.get("url")):
            if link:
                reqs[("link", link.split("?")[0].rstrip("/").lower())] = status
        by_company[nc].append((title, status))
    return reqs, fps, by_company


def do_dedupe(cands, tracker_rows, fuzzy=0.85):
    reqs, fps, by_company = tracker_index(tracker_rows)
    seen_fp, seen_req, out = set(), set(), []
    for j in cands:
        rid = norm_req(j.get("req_id"))
        nc = norm_company(j.get("company"))
        key = (nc, rid)
        fp = j.get("fp") or fingerprint(j.get("company"), j.get("title"), j.get("location"))
        link = (j.get("apply_url") or j.get("url") or "").split("?")[0].rstrip("/").lower()
        dup = None
        if rid and key in reqs:
            dup = f"duplicate (req ID, tracker status {reqs[key] or '?'})"
        elif link and ("link", link) in reqs:
            dup = f"duplicate (same posting link, tracker status {reqs[('link', link)] or '?'})"
        elif fp in fps or fingerprint(j.get("company"), j.get("title"), "") in fps:
            dup = f"duplicate (company+title+city, tracker status {fps.get(fp) or '?'})"
        elif fp in seen_fp or (rid and key in seen_req):
            dup = "duplicate (within this run)"
        else:
            for t, stt in by_company.get(nc, []):
                if title_sim(t, j.get("title", "")) >= fuzzy:
                    dup = f"likely duplicate ('{short(t, 50)}', tracker status {stt or '?'})"
                    break
        seen_fp.add(fp)
        if rid:
            seen_req.add(key)
        if dup:
            j = dict(j)
            j.setdefault("exclude_reasons", []).append(dup)
            j["screen"] = "excluded"
        out.append(j)
    return out


def resolve_to_employer(pool, feed_jobs, min_sim=0.6):
    """Aggregator / LinkedIn finds at companies whose feed was collected: attach the employer posting
    (apply link, req ID, ATS date, description) without opening any page."""
    by_c = defaultdict(list)
    for f in feed_jobs:
        if f.get("source_group") in ("Employer job board",) and f.get("url"):
            by_c[norm_company(f.get("company"))].append(f)
    n = 0
    for j in pool:
        if j.get("apply_url") and not apply_link_problem(j["apply_url"]):
            continue
        cands = by_c.get(norm_company(j.get("company")), [])
        best, bs = None, 0.0
        for f in cands:
            s = title_sim(f.get("title", ""), j.get("title", ""))
            if s > bs and (norm_city(f.get("location")) == norm_city(j.get("location")) or s >= 0.8 or not j.get("location")):
                best, bs = f, s
        if best and bs >= min_sim:
            j["apply_url"] = best["url"]
            j["req_id"] = j.get("req_id") or best.get("req_id", "")
            if best.get("posted_at"):
                j["posted_at_ats"] = best["posted_at"]
            if not j.get("description"):
                j["description"] = best.get("description", "")
            j["resolved_from_feed"] = True
            if j.get("source") == "LinkedIn":
                j["also_on_linkedin"] = "yes"
            n += 1
    return n


# =========================================================================== plan / ledger / ingest
def site_by_name(name):
    for s in CONFIG["sites"]:
        if s["name"].lower() == name.lower():
            return s
    return None


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def render_url(site, program, q, window_h):
    tpl = site.get("url")
    if isinstance(tpl, dict):
        tpl = tpl.get(program) or next(iter(tpl.values()))
    if not tpl:
        return ""
    days = max(1, math.ceil(window_h / 24))
    state = json.dumps({"searchQuery": q, "dateFetchedPastNDays": days, "sortBy": "date"}, separators=(",", ":"))
    return (tpl.replace("{q}", urllib.parse.quote(q)).replace("{seconds}", str(int(window_h * 3600)))
            .replace("{days}", str(days)).replace("{state}", urllib.parse.quote(state))
            .replace("{slug}", slug(q))
            .replace("{gdays}", "today" if window_h <= 24 else "3days" if window_h <= 72 else "week"))


def build_plan(window_h=None, sites=None, skip_sites=None, include_api=True, career_pages=True):
    window_h = window_h or CONFIG["window_hours"]["daily"]
    items = []
    for s in sorted(CONFIG["sites"], key=lambda x: x.get("priority", 9)):
        if sites and s["name"] not in sites:
            continue
        if s.get("enabled") is False and not (sites and s["name"] in sites):
            continue
        if skip_sites and s["name"] in skip_sites:
            continue
        if s["mode"] == "api":
            if include_api:
                items.append({"id": f"{slug(s['name'])}", "site": s["name"], "mode": "api", "program": "+".join(s["programs"]),
                              "family": "", "query": "", "url": "", "status": "todo", "results": 0, "kept": 0, "blocker": "",
                              "priority": s.get("priority", 0)})
            continue
        if s.get("style") == "company":
            if career_pages:
                for c in load_companies():
                    if has_feed(c):
                        continue
                    url = (c.get("search_url") or "").replace("{q}", urllib.parse.quote("data")) or c.get("careers_url", "")
                    items.append({"id": f"career:{slug(c['company'])}", "site": s["name"], "mode": "browser", "program": "professional+newgrad",
                                  "family": "", "query": "data / analyst / machine learning / new grad", "url": url,
                                  "company": c["company"], "status": "todo", "results": 0, "kept": 0, "blocker": "",
                                  "priority": s.get("priority", 1) - (0.1 if c.get("fair_sponsor") == "yes" or c.get("atlanta") == "yes" else 0)})
            continue
        if s.get("style") == "single":
            items.append({"id": f"{slug(s['name'])}", "site": s["name"], "mode": "browser", "program": "+".join(s["programs"]),
                          "family": "", "query": "", "url": render_url(s, s["programs"][0], "", window_h), "status": "todo",
                          "results": 0, "kept": 0, "blocker": "", "priority": s.get("priority", 9)})
            continue
        for prog in s["programs"]:
            for fam, q in CONFIG["query_families"][prog].items():
                if s.get("families") and fam not in s["families"]:
                    continue
                queries = [q["or"]] if s.get("style") == "or" else q["each"]
                for i, qq in enumerate(queries):
                    items.append({"id": f"{slug(s['name'])}:{prog}:{fam}:{i}", "site": s["name"], "mode": "browser",
                                  "program": prog, "family": fam, "query": qq, "url": render_url(s, prog, qq, window_h),
                                  "status": "todo", "results": 0, "kept": 0, "blocker": "", "priority": s.get("priority", 9)})
    ledger = {"created_at": NOW.isoformat(), "run_date": TODAY.isoformat(), "window_hours": window_h, "items": items}
    return ledger


API_LABEL = {"employer-feeds": "Employer feeds", "simplifyjobs-github": "SimplifyJobs GitHub", "speedyapply": "speedyapply", "ask-hn": "Ask HN"}


def apply_collect_meta(ledger, meta):
    """Mark the API items from `collect`'s results: ok -> done, failed -> blocked with the network reason
    (the run then reads those sources another way and says so)."""
    src = (meta or {}).get("sources", {})
    for it in ledger["items"]:
        if it["mode"] != "api":
            continue
        m = src.get(API_LABEL.get(it["id"], it["site"]))
        if not m:
            continue
        it["results"] = m.get("jobs", 0)
        if m.get("status") == "ok":
            it["status"] = "done"
            if it["site"] == "Employer feeds":
                it["note"] = f"{m.get('feeds_ok', 0)}/{m.get('companies_with_feed', 0)} feeds answered; detail pages {m.get('detail_pages_read', 0)}"
        else:
            it["status"], it["blocker"] = "blocked", f"network blocked: {m.get('status')}"
        it["at"] = NOW.isoformat()
    return ledger


def ledger_summary(ledger):
    by = defaultdict(Counter)
    for it in ledger["items"]:
        by[it["site"]][it["status"]] += 1
        by[it["site"]]["results"] += int(it.get("results") or 0)
        by[it["site"]]["kept"] += int(it.get("kept") or 0)
    return by


def est_minutes(items):
    secs = 0
    for it in items:
        if it["mode"] == "browser":
            s = site_by_name(it["site"]) or {}
            secs += s.get("pace_s", 5) + 35  # navigate + extract + ingest call overhead
    return round(secs / 60)


def valid_blocker(text):
    return bool(text) and any(w in text.lower() for w in CONFIG["blocker_words"])


def mark_item(ledger, item_id, status, results=None, kept=None, blocker="", note=""):
    if item_id.endswith("*"):  # 'linkedin:*' marks every LinkedIn item (e.g. one quoted security check)
        hit = [it for it in ledger["items"] if it["id"].startswith(item_id[:-1])]
    else:
        hit = [it for it in ledger["items"] if it["id"] == item_id]
    if not hit:
        raise SystemExit(f"no ledger item {item_id!r} (use `ledger` to list ids)")
    for it in hit:
        it["status"] = status
        if results is not None:
            it["results"] = results
        if kept is not None:
            it["kept"] = kept
        if blocker:
            it["blocker"] = blocker
        if note:
            it["note"] = note
        it["at"] = NOW.isoformat()
    return hit


CARD_SEP = re.compile(r"\s*¦\s*|\t")
SEG_SEP = re.compile(r"\s*·\s*|\s*\n\s*|\s{3,}")
CARD_FLAG_WORDS = ["reposted", "promoted", "easy apply", "viewed", "applied", "actively recruiting", "be an early applicant",
                   "early applicant", "top applicant", "with verification", "verified", "new", "urgently hiring",
                   "responds within", "hiring multiple", "save", "dismiss"]


def parse_card(site, url, title, rest):
    segs = [s.strip() for s in SEG_SEP.split(rest or "") if s.strip()]
    tnorm = norm_title(title)
    flags, age, loc, company, pay, via = [], "", "", "", "", ""
    for s in segs:
        low = s.lower()
        if norm_title(s) == tnorm or low.endswith("with verification") and norm_title(s.replace("with verification", "")) == tnorm:
            continue
        if low.startswith("via "):
            via = s[4:].strip()
            continue
        if "reposted" in low:
            flags.append("Reposted")
        if any(low == w or low.startswith(w) for w in CARD_FLAG_WORDS) and parse_age(s) is None:
            flags.append(s)
            continue
        if not age and parse_age(s) is not None and len(s) < 40:
            age = s
            continue
        if not pay and "$" in s:
            pay = s
            continue
        if not loc and (is_us(s) != "unknown" or re.search(r"\bremote\b|\bhybrid\b|\bon-?site\b", low)) and len(s) < 80:
            loc = s
            continue
        if not company and len(s) < 70 and not re.search(r"\d+ applicants?|alumni|connections?|benefits|school", low):
            company = s
    return {"url": url, "title": title, "company": company, "location": loc, "age_text": age, "pay": pay, "via": via,
            "card_flags": flags, "reposted": "Reposted" in flags, "card_text": short(rest, 400)}


def parse_cards(text, site):
    """Lines 'url ¦ title ¦ card text' (from extractors.js), JSON lines, or a JSON array."""
    t = text.strip()
    out = []
    if t.startswith("["):
        for r in json.loads(t):
            out.append(r if "card_text" in r else {**parse_card(site, r.get("url", ""), r.get("title", ""), r.get("text", "")), **r})
        return out
    for line in t.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("{"):
            r = json.loads(line)
            out.append(r if "card_text" in r else {**parse_card(site, r.get("url", ""), r.get("title", ""), r.get("text", "")), **r})
            continue
        parts = CARD_SEP.split(line, maxsplit=2)
        if len(parts) < 2:
            continue
        url, title = parts[0], parts[1]
        rest = parts[2] if len(parts) > 2 else ""
        if re.match(r"^(?:li|in):\w+$", url):
            url = canon_url(site, url)
        if not url.startswith("http"):  # 'title ¦ company ¦ ...' without a URL
            url, title, rest = "", parts[0], " · ".join(parts[1:])
        out.append(parse_card(site, url, title, rest))
    return out


def canon_url(site, url):
    u = (url or "").strip()
    if u.startswith("li:"):
        return f"https://www.linkedin.com/jobs/view/{u[3:]}/"
    if u.startswith("in:"):
        return f"https://www.indeed.com/viewjob?jk={u[3:]}"
    url = u
    m = LI_ID_RX.search(url or "")
    if m:
        return f"https://www.linkedin.com/jobs/view/{m.group(1)}/"
    if "indeed.com" in (url or ""):
        jk = re.search(r"jk=([0-9a-f]{12,})", url)
        if jk:
            return f"https://www.indeed.com/viewjob?jk={jk.group(1)}"
    return (url or "").split("#")[0]


def do_ingest(ledger, item_id, text, pool_path, site=None, extra=None):
    items = [it for it in ledger["items"] if it["id"] == item_id]
    if not items:
        raise SystemExit(f"no ledger item {item_id!r}")
    it = items[0]
    site = site or it["site"]
    cards = parse_cards(text, site)
    rows, seen = [], set()
    for c in cards:
        if not c.get("title"):
            continue
        c["url"] = canon_url(site, c.get("url", ""))
        k = c["url"] or (c.get("company", "") + "|" + c["title"])
        if k in seen:
            continue
        seen.add(k)
        src_group = site if site in CONTROLLED["source_group"] else {"Career pages": "Company website", "Simplify FAANG+": "Simplify Jobs"}.get(site, "Other")
        row = {"company": c.get("company") or it.get("company", ""), "title": c["title"], "location": c.get("location", ""),
               "url": c["url"], "apply_url": "" if (c["url"] and apply_link_problem(c["url"])) else c["url"],
               "age_text": c.get("age_text", ""), "pay": c.get("pay", ""), "card_flags": c.get("card_flags", []),
               "reposted": c.get("reposted", False), "snippet": c.get("card_text", ""), "source": site,
               "source_group": src_group, "ledger_item": it["id"], "program_hint": "newgrad" if it.get("program") == "newgrad" else "",
               "found_at": NOW.isoformat()}
        if "newgrad" in str(it.get("program")) and site in ("Simplify FAANG+", "newgrad-jobs.com"):
            row["program_hint"] = "newgrad"
        cr = company_row(row["company"])
        if cr:
            for k2 in ("sector", "startup", "atlanta", "gt_fair", "fair_sponsor"):
                row[k2] = cr.get(k2, "")
        role = (site_by_name(site) or {}).get("role", "discovery")
        if cr and role in ("discovery", "graduate") and CONFIG.get("discovery", {}).get("drop_known_companies", True):
            row["known_company"] = cr["company"]  # harvested directly from its own board / career page
        if c.get("via"):
            row["via"] = c["via"]
        if extra:
            row.update(extra)
        rows.append(row)
    append_jsonl(rows, pool_path)
    it["status"], it["results"], it["at"] = "done", it.get("results", 0) + len(rows), NOW.isoformat()
    return rows


# =========================================================================== shortlist and finalize
def build_shortlist(screened, tracker_rows, feed_jobs=None, threshold=None):
    threshold = RUBRIC["claude_threshold"] if threshold is None else threshold
    pool = linkedin_age_check([dict(j) for j in screened])
    if feed_jobs:
        resolve_to_employer([j for j in pool if j.get("screen") == "pass"], feed_jobs)
    pool = do_dedupe(pool, tracker_rows)
    passed = [j for j in pool if j.get("screen") == "pass"]
    # one row per role: an employer-feed copy beats an aggregator copy of the same fingerprint
    best = {}
    for j in passed:
        k = j["fp"]
        cur = best.get(k)
        rank = (0 if j.get("source_group") == "Employer job board" else 1, -(j.get("pre_score") or 0))
        if not cur or rank < cur[0]:
            if cur:
                j.setdefault("also_seen_on", []).extend([cur[1].get("source")] + cur[1].get("also_seen_on", []))
            best[k] = (rank, j)
        else:
            cur[1].setdefault("also_seen_on", []).append(j.get("source"))
    rows = [b[1] for b in best.values()]
    for j in rows:
        if "LinkedIn" in ([j.get("source")] + j.get("also_seen_on", [])):
            j["also_on_linkedin"] = "yes"
        j["needs"] = needs_for(j)
    short_ = sorted([j for j in rows if (j.get("pre_score") or 0) >= threshold], key=lambda j: -(j.get("pre_score") or 0))
    below = [j for j in rows if (j.get("pre_score") or 0) < threshold]
    return pool, short_, below


def needs_for(j):
    n = []
    if not j.get("apply_url") or apply_link_problem(j.get("apply_url")):
        n.append("employer link")
    if j.get("age_hours") is None:
        n.append("date")
    if not (j.get("description") or "").strip():
        n.append("posting text")
    if j.get("source") == "LinkedIn" and not (j.get("resolved_from_feed") or j.get("opened")):
        n.append("open: Reposted?")
    return n


CONTROLLED = {
    "track_group": ["Data science", "ML / AI engineering", "Data / analytics", "Analytics & consulting",
                    "Risk & credit analytics", "Model risk", "Quant", "Forward deployed & solutions", "Data engineering",
                    "Finance & investment analytics"],
    "sector_group": ["Fintech", "Finance & banking", "Tech / AI", "Consulting", "Real estate tech", "Other"],
    "source_group": ["LinkedIn", "HiringCafe", "Built In", "Indeed", "Jobright", "Simplify Jobs", "TrueUp", "Levels.fyi",
                     "Welcome to the Jungle", "eFinancialCareers", "Y Combinator", "Wellfound", "Employer job board",
                     "Company website", "Ask HN", "Interstride", "GT CareerBuzz", "Other"],
    "priority_group": ["High", "Medium", "Stretch", "Blocked"],
    "also_on_linkedin": ["yes", "no", "unknown"],
    "status": ["Not applied", "Applied", "Recruiter contacted", "Interview scheduled", "Interview completed",
               "Offer received", "Rejected", "Closed", "Skipped"],
}
REQUIRED_MAIN = ["company", "title", "location", "apply_link", "score", "priority_group", "track_group", "sector_group",
                 "source_group", "also_on_linkedin", "status", "ev_experience", "batch", "posted", "fp"]
REQUIRED_NEWGRAD = ["company", "title", "category", "location", "source", "open_status", "quick_fit", "apply_link",
                    "date_checked", "status", "posted", "fp", "score"]


def band(score):
    b = RUBRIC["bands"]
    return "High" if score >= b["High"] else "Medium" if score >= b["Medium"] else "Stretch" if score >= b["Stretch"] else "Below"


def posted_label(j):
    iso_s = j.get("posted_at_ats") or j.get("posted_at")
    if not iso_s:
        return ""
    d = dt.datetime.fromisoformat(iso_s).astimezone(dt.timezone(dt.timedelta(hours=-4))).date()
    if j.get("posted_at_ats") or (j.get("source_group") == "Employer job board" and not j.get("age_text")):
        src = "employer page / ATS"
    elif j.get("age_text"):
        src = f"{j['age_text']}, {j.get('source')}"
    else:
        src = j.get("source", "")
    return f"{d:%b %d} ({src})" if src else f"{d:%b %d}"


def to_row(j, kind):
    """A scored candidate -> a tracker row. Claude's `score` block overrides the pre-score per slot."""
    prog = "newgrad" if kind == "newgrad" else "professional"
    w = RUBRIC[prog]
    exp_key = "experience" if prog == "professional" else "eligibility"
    slots = ["skills", exp_key, "role_fit", "domain_quals", "preferences"]
    sc = j.get("score") if isinstance(j.get("score"), dict) else {}
    pts, evs = {}, {}
    for s in slots:
        v = sc.get(s)
        auto = (j.get("pre_evidence") or {}).get(s, "")
        if isinstance(v, dict):
            pts[s], evs[s] = v.get("pts"), v.get("ev") or auto
        elif v is not None:
            pts[s], evs[s] = v, auto
        else:
            pts[s] = (j.get("pre_components") or {}).get(s)
            evs[s] = (j.get("pre_evidence") or {}).get(s, "") + " (pre-score)"
        pts[s] = max(0, min(w[s], int(round(float(pts[s] or 0)))))
    if prog == "professional" and j.get("ev_experience_auto") and not evs["experience"].startswith("Asks"):
        evs["experience"] = (j["ev_experience_auto"] + " " + evs["experience"]).strip()  # find-jobs.md: always "Asks: … Has: …"
    total = sum(pts.values())
    pri = band(total)
    if prog == "professional" and pri == "High" and pts["experience"] < 15:
        pri = "Medium"  # High also needs the years to fit (find-jobs.md ranking rule)
    link = j.get("apply_url") or j.get("apply_link") or ""
    batch = j.get("batch") or f"Daily {TODAY.isoformat()}"
    base = {"id": f"d{TODAY:%m%d}-{j['fp']}", "fp": j["fp"], "company": j.get("company", ""), "title": j.get("title", ""),
            "location": j.get("location", ""), "apply_link": link, "req_id": j.get("req_id", ""), "posted": j.get("posted") or posted_label(j),
            "source": j.get("source", ""), "date_checked": TODAY.isoformat(), "status": "Not applied", "batch": batch,
            "score": total, "priority": pri, "h1b_history": j.get("h1b_history", ""), "flags": "; ".join(j.get("flags", [])),
            "notes": j.get("notes", ""), "program": prog, "_new": True,
            "score_max": {CONFIG["rubric"]["slot_fields"][s]: w[s] for s in slots}}
    for s in slots:
        base[RUBRIC["slot_fields"][s]] = pts[s]
        base[RUBRIC["slot_evidence"][s]] = evs[s]
    if kind == "newgrad":
        base.update({"category": NG["category"].get(j.get("family"), "SWE/SDE"),
                     "quick_fit": {"High": "High", "Medium": "Medium"}.get(pri, "Low"),
                     "open_status": j.get("open_status", "Open"), "deadline": j.get("deadline", ""),
                     "deadline_date": j.get("deadline_date", "")})
        return base
    base.update({"priority_group": pri if pri != "Below" else "Stretch",
                 "track": j.get("track") or PROF["track_group"].get(j.get("family"), ""),
                 "track_group": j.get("track_group") or track_group_for(j),
                 "sector": j.get("sector", ""), "sector_group": CONFIG["sectors"]["sector_group"].get(j.get("sector", ""), "Other"),
                 "source_group": j.get("source_group") if j.get("source_group") in CONTROLLED["source_group"] else "Other",
                 "also_on_linkedin": j.get("also_on_linkedin") or ("yes" if j.get("source") == "LinkedIn" else "unknown"),
                 "verification": "Verified", "arrangement": j.get("arrangement", ""), "employment": j.get("employment_type") or "Full-time",
                 "experience": j.get("experience") or (f"{j['years_required']}+ yrs" if j.get("years_required") else "Not stated"),
                 "salary": j.get("salary") or j.get("pay") or "Not listed", "key_skills": j.get("key_skills", ""),
                 "gaps": j.get("gaps", "") or ("Missing from resume: " + ", ".join((j.get("jd_match") or {}).get("missing", [])[:6]) if (j.get("jd_match") or {}).get("missing") else ""),
                 "disqualifiers": j.get("disqualifiers", ""), "resume": j.get("resume") or j.get("resume_suggested", ""),
                 "why": j.get("why", ""), "lead_with": j.get("lead_with", ""), "verify": j.get("verify", "")})
    if j.get("startup") == "yes" and "startup" not in base["notes"].lower():
        base["notes"] = ("Startup. " + base["notes"]).strip()
    return base


def track_group_for(j):
    fam = j.get("family")
    t = (j.get("title") or "").lower()
    if fam == "risk_quant":
        if "model risk" in t or "model validation" in t or "validator" in t:
            return "Model risk"
        if "quant" in t:
            return "Quant"
        return "Risk & credit analytics"
    return PROF["track_group"].get(fam, "Data science")


def validate_rows(rows, kind="main", window_h=None):
    window_h = window_h or CONFIG["window_hours"]["widened"]
    problems = []
    req = REQUIRED_NEWGRAD if kind == "newgrad" else REQUIRED_MAIN
    for i, r in enumerate(rows):
        name = f"row {i + 1} ({r.get('company', '?')} / {short(r.get('title', '?'), 50)})"
        problems += [f"{name}: missing {f}" for f in req if str(r.get(f, "")).strip() == ""]
        if kind == "newgrad":
            if r.get("quick_fit") not in (None, "", "High", "Medium", "Low"):
                problems.append(f"{name}: quick_fit must be High / Medium / Low")
            if r.get("open_status") not in (None, "", "Open", "Opening soon", "Closed"):
                problems.append(f"{name}: open_status must be Open / Opening soon / Closed")
        else:
            for f, allowed in CONTROLLED.items():
                v = r.get(f)
                if v not in (None, "") and v not in allowed:
                    problems.append(f"{name}: {f}={v!r} not in {allowed}")
            if r.get("priority_group") == "High" and float(r.get("score") or 0) < RUBRIC["bands"]["High"]:
                problems.append(f"{name}: High priority but fit {r.get('score')} < {RUBRIC['bands']['High']}")
        try:
            s = float(r.get("score", 0) or 0)
            if r.get("_new") and s < RUBRIC["add_threshold"]:
                problems.append(f"{name}: score {s:g} below the add threshold {RUBRIC['add_threshold']} — count it, don't add it")
        except (TypeError, ValueError):
            problems.append(f"{name}: score not a number")
        lp = apply_link_problem(r.get("apply_link"))
        if lp:
            problems.append(f"{name}: apply_link {lp} — find the employer's own posting or do not add the role")
        ver = str(r.get("verification", "")).lower()
        if "unverified" in ver or "list only" in ver:
            problems.append(f"{name}: verification '{r.get('verification')}' — new rows need a date and employer page (rule 18)")
        if re.search(r"\brepost", str(r.get("flags", "")), re.I):
            problems.append(f"{name}: marked reposted — reposted roles are excluded, not added")
        if r.get("_new") and r.get("status") not in (None, "", "Not applied"):
            problems.append(f"{name}: new rows must start as 'Not applied'")
        if r.get("_new") and kind == "main" and not r.get("posted"):
            problems.append(f"{name}: no posting date or age (rule 18)")
    return problems


def do_finalize(scored, main_rows, ng_rows):
    """scored: shortlist items with Claude's `score` blocks (and `verified: true`), or `drop` reasons.
    -> new rows per tracker (validated), merged trackers, and a funnel of what was not added."""
    new_main, new_ng, not_added = [], [], Counter()
    leads = []
    existing = list(main_rows) + list(ng_rows)
    for j in scored:
        if j.get("drop"):
            not_added[j["drop"]] += 1
            if j.get("lead"):
                leads.append({"company": j.get("company"), "title": j.get("title"), "seen_on": j.get("source"), "why": j["drop"]})
            continue
        if j.get("screen") == "excluded":
            not_added[re.sub(r" \(.*|:.*", "", (j.get("exclude_reasons") or ["excluded"])[0])] += 1
            continue
        if not j.get("score"):
            not_added["not scored by Claude"] += 1
            continue
        kind = "newgrad" if j.get("tracker") == "newgrad" or j.get("program") == "newgrad" else "main"
        row = to_row(j, kind)
        if row["score"] < RUBRIC["add_threshold"]:
            not_added[f"score below {RUBRIC['add_threshold']}"] += 1
            continue
        probs = validate_rows([row], kind)
        if probs:
            not_added["failed validation"] += 1
            leads.append({"company": row["company"], "title": row["title"], "seen_on": j.get("source"), "why": "; ".join(p.split(": ", 1)[1] for p in probs)})
            continue
        if do_dedupe([{"company": row["company"], "title": row["title"], "location": row["location"], "req_id": row["req_id"],
                       "fp": row["fp"], "apply_url": row["apply_link"]}], existing + new_main + new_ng)[0].get("screen") == "excluded":
            not_added["duplicate at finalize"] += 1
            continue
        (new_ng if kind == "newgrad" else new_main).append(row)
    for r in new_main + new_ng:
        r.pop("_new", None)
    return new_main, new_ng, not_added, leads



REQ_LINE = re.compile(r"years?|yrs|experience|required|requirement|qualification|must|degree|bachelor|master|ph\.?d|"
                      r"sponsor|visa|authoriz|citizen|clearance|salary|\$\d|pay range|compensation|hybrid|remote|on-?site|"
                      r"python|sql|machine learning|model|statistic|risk|credit|cfa|you have|you will|responsib", re.I)


def brief_text(desc, limit=1100):
    """The lines of a posting that decide gates and score (years, degrees, must-haves, sponsorship, pay)."""
    parts = [p.strip() for p in re.split(r"\n|(?<=[.;])\s+(?=[A-Z•\-*])|\s[•·]\s", desc or "") if p.strip()]
    keep, n = [], 0
    for p_ in parts:
        if REQ_LINE.search(p_) and len(p_) > 12:
            p_ = short(p_, 220)
            if p_ in keep:
                continue
            keep.append(p_)
            n += len(p_) + 3
            if n > limit:
                break
    return " | ".join(keep)


def merge_scores(shortlist, score_lines):
    """scores.jsonl: one {fp, score:{slot:{pts, ev}}, resume, key_skills, gaps, why, lead_with, verify,
    apply_url?, posted?, description?, drop?, lead?} per line -> shortlist items with Claude's fields."""
    by = {}
    for r in score_lines:
        by.setdefault(r.get("fp"), {}).update(r)
    out = []
    for j in shortlist:
        if j["fp"] in by:
            j = {**j, **{k: v for k, v in by[j["fp"]].items() if v not in (None, "")}}
            out.append(j)
    missing = [fp for fp in by if fp not in {j["fp"] for j in shortlist}]
    return out, missing


# =========================================================================== click-to-update trackers (artifact database)
USER_FIELDS = {"status", "application_date", "reject_reason", "reject_date", "reject_sponsorship", "referral_status",
               "referral_contact", "referral_date", "my_note", "interview_date", "contact", "follow_up", "updated_at", "updated_by"}
DOC_ID_RX = re.compile(r"[^A-Za-z0-9_\-.~:@+]")


def doc_id(r):
    raw = str(r.get("id") or r.get("_id") or r.get("fp") or fingerprint(r.get("company"), r.get("title"), r.get("location")))
    return DOC_ID_RX.sub("-", raw)[:150].strip(".") or "row"


def row_to_doc(r):
    d = {k: v for k, v in r.items() if not k.startswith("_") and v not in (None, "")}
    d.setdefault("added_at", r.get("added_at") or NOW.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
    if isinstance(d.get("flags"), list):
        d["flags"] = "; ".join(d["flags"])
    return d


def tracker_docs(rows, existing_ids, outdir, chunk=50):
    """Per-document JSON files + batch manifests for ArtifactData `batch` (each entry names a file_path, so the
    tool call stays small). Only new rows are written (op set, no version needed); ids already in the tracker are
    skipped, because the page owns those docs once she has touched them. To correct one, ArtifactData get it and
    update it pinned to the version that get shows."""
    outdir = Path(outdir)
    (outdir / "docs").mkdir(parents=True, exist_ok=True)
    writes, seen = [], set()
    for r in rows:
        i = doc_id(r)
        if i in existing_ids or i in seen:
            continue  # already in the tracker: her edits live there; an update would need the doc's version
        seen.add(i)
        d = {k: v for k, v in row_to_doc(r).items() if k not in USER_FIELDS or k == "status" and v == "Not applied"}
        fp_ = outdir / "docs" / f"{i}.json"
        dump_json(d, fp_)
        writes.append({"op": "set", "collection": "roles", "doc_id": i, "file_path": str(fp_.resolve())})
    batches = [writes[k:k + chunk] for k in range(0, len(writes), chunk)]
    for n, b in enumerate(batches, 1):
        dump_json(b, outdir / f"batch_{n:03d}.json")
    return batches


def tracker_export(folder):
    """ArtifactData `list` (collection roles, out_dir=folder) -> rows; each doc file is {id, data, version}."""
    rows = []
    for p_ in sorted(Path(folder).rglob("*.json")):
        try:
            d = json.loads(p_.read_text(encoding="utf-8"))
        except Exception:
            continue
        body = d.get("data") if isinstance(d, dict) and isinstance(d.get("data"), dict) else d
        if not isinstance(body, dict) or not body.get("title"):
            continue
        body = dict(body)
        body["_id"] = d.get("id") or p_.stem
        body["_version"] = d.get("version")
        rows.append(body)
    return rows


def state_unpack(folder, keep_days=60):
    """runstate docs (one per run, create-only, in the professional tracker's database) -> {"seen": {fp: iso}}.
    One doc per run keeps every write a create (no versions) and each doc far below the 256 KB limit."""
    seen, docs = {}, 0
    cutoff = (NOW - dt.timedelta(days=keep_days)).isoformat()
    for p_ in sorted(Path(folder).rglob("*.json")):
        d = load_json(p_, default={})
        d = d.get("data", d) if isinstance(d, dict) else {}
        if not isinstance(d.get("seen"), dict):
            continue
        docs += 1
        for fp_, when in d["seen"].items():
            if when > cutoff and (fp_ not in seen or when < seen[fp_]):
                seen[fp_] = when
    return {"seen": seen, "docs": docs, "updated_at": NOW.isoformat()}


def state_pack(state, since):
    """Entries first seen at or after `since` (this run) -> (doc id, doc, count). Dates are kept to the minute."""
    new = {k: v[:16] for k, v in (state.get("seen") or {}).items() if v >= since}
    did = "run-" + NOW.strftime("%Y%m%dT%H%M")
    return did, {"run_at": NOW.isoformat(timespec="minutes"), "seen": new}, len(new)


def stale_runstate(folder, keep_days=60):
    """runstate doc ids older than keep_days (delete them, pinned to version 1: runstate docs are never edited)."""
    cutoff = (NOW - dt.timedelta(days=keep_days)).strftime("run-%Y%m%dT%H%M")
    return sorted(p_.stem for p_ in Path(folder).rglob("run-*.json") if p_.stem < cutoff)


SCHOOLS = ["Georgia Tech", "EDHEC", "Sun Yat-sen University"]


def referral_pack(company, rows):
    q = urllib.parse.quote
    links = [{"label": f"LinkedIn: {s} alumni at {company}", "url": f"https://www.linkedin.com/search/results/people/?keywords={q(company + ' ' + s)}"} for s in SCHOOLS[:2]]
    links += [{"label": "LinkedIn: recruiters", "url": f"https://www.linkedin.com/search/results/people/?keywords={q(company + ' recruiter data')}"},
              {"label": "1point3acres 内推", "url": f"https://www.google.com/search?q={q('site:1point3acres.com 内推 ' + company)}"},
              {"label": "Blind referral threads", "url": f"https://www.google.com/search?q={q('site:teamblind.com ' + company + ' referral')}"}]
    top = max(rows, key=lambda r: float(r.get("score") or 0))
    skills = [x.strip() for x in re.split(r"[;,]|\(|\)", top.get("ev_skills") or "") if x.strip() and not x.strip().startswith(("has ", "missing"))][:3]
    title = top.get("title", "")
    others = [r.get("title", "") for r in rows if r is not top][:2]
    # Every fact below is on her resumes (resumes/*.md); keep it that way when editing.
    if any((r.get("program") == "newgrad") for r in rows):
        body = (". I build ML systems end to end: as a data science intern at TransUnion I built a production income-prediction pipeline "
                "(a 5,000+ line R package adopted across teams) over 120M+ records, and at Georgia Tech I built an LLM pipeline with FAISS-based "
                "retrieval, LoRA fine-tuning and GRPO that reached 28.7% on SWE-bench Verified.")
    else:
        body = (". My background is ML modeling on large financial data: as a data science intern at TransUnion I built credit-risk and "
                "income models across 120M+ records, and before grad school I spent four years building investment models behind $5.5B+ "
                "of real estate investment decisions.")
    draft = (f"Hi <name>, I'm Yuanting, an M.S. Computational Science & Engineering graduate from Georgia Tech (May 2026) with a "
             f"financial engineering master's and CFA. I'm applying for the {title} role at {company}"
             + (f" (and {', '.join(others)})" if others else "")
             + body
             + (f" The posting's focus on {', '.join(skills)} matches that work closely." if skills else "")
             + f" Would you be open to referring me, or to a 15-minute chat about the team? The posting: {top.get('apply_link', '')} Thank you!")
    return links, draft


def pack_id(company):
    return DOC_ID_RX.sub("-", norm_company(company).replace(" ", "-"))[:150].strip(".-") or "company"


def do_referrals(rows, existing_packs=()):
    """One referral pack per company with a High, open role, written to the `referrals` collection (doc id =
    pack_id). Packs are create-only: a company that already has one is skipped, so no write ever needs the
    version of a document the page may have changed. Role docs are never touched; the page joins packs to roles
    by company name (pack['companies'] lists every display name seen). The draft uses only resume facts."""
    by = defaultdict(list)
    for r in rows:
        p_ = r.get("priority_group") or r.get("priority") or r.get("quick_fit")
        if r.get("gate_excluded"):
            continue
        if p_ == "High" and (r.get("status") or "Not applied") in ("Not applied", "") and r.get("referral_status") not in ("Referred", "No referral", "Not needed"):
            by[pack_id(r.get("company"))].append(r)
    have = set(existing_packs)
    packs = []
    for k, rs in by.items():
        if k in have:
            continue
        rs = sorted(rs, key=lambda r: -float(r.get("score") or 0))
        links, draft = referral_pack(rs[0].get("company", ""), rs)
        packs.append({"_id": k, "company": rs[0].get("company", ""), "companies": sorted({r.get("company", "") for r in rs if r.get("company")}),
                      "links": links, "draft": draft, "top_title": rs[0].get("title", ""), "top_score": rs[0].get("score"),
                      "created_at": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")})
    return packs

# =========================================================================== audit, report, stats, calibrate
def do_audit_ledger(ledger):
    rows, ok = [], True
    for site, c in ledger_summary(ledger).items():
        todo, done, blocked = c.get("todo", 0), c.get("done", 0), c.get("blocked", 0)
        bad_block = [it for it in ledger["items"] if it["site"] == site and it["status"] == "blocked" and not valid_blocker(it.get("blocker", ""))]
        status = "OK" if todo == 0 and not bad_block else ("BLOCKED (quoted)" if todo == 0 and blocked else "SHORT")
        if bad_block:
            status = "SHORT (blocker not an allowed reason)"
        if todo or bad_block:
            ok = False
        if todo == 0 and blocked and not bad_block and not done:
            status = "BLOCKED (quoted)"
        quotes = "; ".join(dict.fromkeys(it.get("blocker", "") for it in ledger["items"] if it["site"] == site and it.get("blocker")))
        rows.append({"site": site, "done": done, "blocked": blocked, "todo": todo, "results": c.get("results", 0),
                     "kept": c.get("kept", 0), "status": status, "blocker": short(quotes, 70)})
    return ok, rows


def do_audit_legacy(run, log_rows):
    mins = dict(CONFIG["minimums_legacy"]["all"])
    by = defaultdict(list)
    for r in log_rows:
        by[r.get("site", "").strip().lower()].append(r)
    rows, ok = [], True
    for site, need in mins.items():
        rs = by.get(site.lower(), [])
        q = sum(int(r.get("queries") or 0) for r in rs)
        blocker = "; ".join(r.get("blocker", "") for r in rs if r.get("blocker"))
        if need == "all":
            total = sum(int(r.get("companies_total") or 0) for r in rs)
            done = sum(int(r.get("companies_done") or r.get("queries") or 0) for r in rs)
            passed, need_s = bool(rs) and (total == 0 or done >= total), f"{done}/{total or '?'} companies"
        else:
            passed, need_s = q >= int(need), f"{q}/{need}"
        status = "OK" if passed else ("BLOCKED (quoted)" if valid_blocker(blocker) else "SHORT")
        ok = ok and status != "SHORT"
        rows.append({"site": site, "queries": need_s, "status": status, "blocker": blocker})
    return ok, rows


def funnel(jobs):
    c = Counter()
    for j in jobs:
        if j.get("screen") == "pass":
            c["passed gates"] += 1
        for r in j.get("exclude_reasons", []):
            c[re.sub(r" \(.*|:.*", "", r)] += 1
    return dict(c.most_common())


def do_stats(rows):
    src = Counter(r.get("source_group") or r.get("source") or "?" for r in rows)
    li_only = sum(1 for r in rows if (r.get("source_group") or r.get("source") or "") == "LinkedIn")
    li_any = sum(1 for r in rows if (r.get("also_on_linkedin") or "").lower() == "yes" or r.get("source_group") == "LinkedIn")
    by_batch = defaultdict(Counter)
    for r in rows:
        by_batch[r.get("batch") or "?"][r.get("priority_group") or r.get("priority") or r.get("quick_fit") or "?"] += 1
    return {"rows": len(rows), "by_source": dict(src.most_common()), "linkedin_only": li_only,
            "linkedin_only_share": round(li_only / max(1, len(rows)), 3), "on_linkedin_any": li_any,
            "by_priority": dict(Counter(r.get("priority_group") or r.get("priority") or r.get("quick_fit") or "?" for r in rows)),
            "by_status": dict(Counter(r.get("status") or "?" for r in rows)),
            "by_batch": {b: dict(c) for b, c in sorted(by_batch.items())}}


def do_calibrate(rows):
    """Does the score predict what she does? Apply rate by band and by score bucket, and AUC of score vs applied."""
    acted = [r for r in rows if r.get("status") in ("Applied", "Rejected", "Interview scheduled", "Interview completed",
                                                    "Offer received", "Recruiter contacted", "Skipped")]
    def applied(r):
        return r.get("status") != "Skipped"
    def score(r):
        try:
            return float(r.get("score") or 0)
        except (TypeError, ValueError):
            return 0.0
    pos = [score(r) for r in acted if applied(r) and score(r) > 0]
    neg = [score(r) for r in acted if not applied(r) and score(r) > 0]
    auc = None
    if pos and neg:
        wins = sum((p > n) + 0.5 * (p == n) for p in pos for n in neg)
        auc = round(wins / (len(pos) * len(neg)), 3)
    by_band = defaultdict(Counter)
    for r in acted:
        by_band[r.get("priority_group") or r.get("priority") or "?"]["applied" if applied(r) else "skipped"] += 1
    buckets = defaultdict(Counter)
    for r in acted:
        s = score(r)
        if s:
            buckets[f"{int(s // 5 * 5)}-{int(s // 5 * 5) + 4}"]["applied" if applied(r) else "skipped"] += 1
    reasons = Counter()
    for r in acted:
        if not applied(r):
            n = (r.get("notes") or "").lower()
            for k, lab in (("repost", "reposted"), ("too senior", "too senior"), ("gap too large", "too senior"), ("experience gap", "too senior"),
                           ("sponsorship", "sponsorship"), ("not found", "no employer posting"), ("no longer", "closed"),
                           ("duplicate", "duplicate"), ("principal", "too senior"), ("too advanced", "too senior")):
                if k in n:
                    reasons[lab] += 1
                    break
            else:
                reasons["not stated"] += 1
    return {"acted_on": len(acted), "applied": len(pos), "skipped": len(neg), "auc_score_vs_applied": auc,
            "mean_score_applied": round(sum(pos) / max(1, len(pos)), 1), "mean_score_skipped": round(sum(neg) / max(1, len(neg)), 1),
            "by_band": {k: dict(v) for k, v in by_band.items()}, "by_score_bucket": {k: dict(v) for k, v in sorted(buckets.items())},
            "skip_reasons": dict(reasons.most_common()),
            "reading": "AUC 0.5 = the score does not predict applying; 0.7+ = useful. Skip reasons that are gates "
                       "(reposted, too senior, sponsorship, closed) belong in screening, not in the score."}


def build_report(ledger, screened_funnel, new_main, new_ng, not_added, leads, collect_meta=None, prefilled=None):
    ok, arows = do_audit_ledger(ledger)
    lines = [f"# Daily job run {TODAY.isoformat()}" + ("" if ok else " — INCOMPLETE"), ""]
    lines.append("## Coverage")
    lines.append("| site | done | blocked | todo | results | status |")
    lines.append("|---|---|---|---|---|---|")
    for r in arows:
        lines.append(f"| {r['site']} | {r['done']} | {r['blocked']} | {r['todo']} | {r['results']} | {r['status']}{(': ' + r['blocker']) if r['blocker'] else ''} |")
    if collect_meta:
        src = collect_meta.get("sources", {})
        lines += ["", "Machine sources: " + "; ".join(f"{k}: {v.get('status')} ({v.get('jobs', 0)} jobs)" for k, v in src.items())]
        if collect_meta.get("feeds_failed"):
            lines.append("Feeds that failed: " + ", ".join(f"{f['company']} ({f['status']})" for f in collect_meta["feeds_failed"][:25]))
    lines += ["", "## Funnel", ", ".join(f"{k} {v}" for k, v in (screened_funnel or {}).items())]
    if not_added:
        lines.append("Not added after scoring: " + ", ".join(f"{k} {v}" for k, v in not_added.items()))
    for label, rows in (("Professional tracker", new_main), ("New-grad tracker", new_ng)):
        lines += ["", f"## {label}: {len(rows)} added"]
        for r in sorted(rows, key=lambda x: -x["score"])[:40]:
            lines.append(f"- {r['score']} {r['priority']} · {r['company']} — {r['title']} · {r.get('location', '')} · {r.get('posted', '')}")
    if leads:
        lines += ["", f"## Leads not added ({len(leads)})"] + [f"- {l['company']} — {l['title']} (seen on {l.get('seen_on')}): {l['why']}" for l in leads[:30]]
    if prefilled:
        lines += ["", "## Pre-filled (not submitted)"] + [f"- {p}" for p in prefilled]
    return "\n".join(lines) + "\n", ok



# =========================================================================== DOL LCA sponsors (registry growth + sponsorship evidence)
SP_CFG = CONFIG.get("sponsors", {})
SOC_GROUP = {code: g for g, codes in SP_CFG.get("soc_groups", {}).items() for code in codes}


def iter_xlsx_rows(path):
    """Rows of the first sheet of an .xlsx as dicts. openpyxl (read-only) when installed, else a stdlib
    streaming parser (zipfile + iterparse)."""
    wb = None
    try:
        import openpyxl  # noqa: F401  (GitHub Actions: pip install openpyxl)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:  # not installed, or a file openpyxl can't open: use the stdlib reader
        wb = None
    if wb is not None:
        it = wb.worksheets[0].iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(it)]
        for r in it:
            yield {h: ("" if v is None else v) for h, v in zip(head, r)}
        return
    import zipfile
    import xml.etree.ElementTree as ET
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    z = zipfile.ZipFile(path)
    shared = []
    if "xl/sharedStrings.xml" in z.namelist():
        for _, el in ET.iterparse(z.open("xl/sharedStrings.xml")):
            if el.tag == ns + "si":
                shared.append("".join(t.text or "" for t in el.iter(ns + "t")))
                el.clear()
    sheet = next(n for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
    head = None

    def col(ref):
        n = 0
        for ch in re.match(r"[A-Z]+", ref).group(0):
            n = n * 26 + ord(ch) - 64
        return n - 1
    for _, el in ET.iterparse(z.open(sheet)):
        if el.tag != ns + "row":
            continue
        vals = {}
        for cel in el.iter(ns + "c"):
            v = cel.find(ns + "v")
            txt = v.text if v is not None else "".join(t.text or "" for t in cel.iter(ns + "t"))
            if cel.get("t") == "s" and txt is not None:
                txt = shared[int(txt)]
            vals[col(cel.get("r"))] = txt or ""
        el.clear()
        row = [vals.get(i, "") for i in range(max(vals) + 1)] if vals else []
        if head is None:
            head = [str(h).strip() for h in row]
            continue
        yield dict(zip(head, row))


def annual_wage(v, unit):
    try:
        x = float(str(v).replace("$", "").replace(",", "").split("-")[0])
    except ValueError:
        return None
    mult = {"hour": 2080, "week": 52, "bi-weekly": 26, "month": 12, "year": 1}.get(str(unit).strip().lower(), 1)
    return x * mult


def build_sponsors(paths):
    """LCA rows -> {norm employer: {name, total, groups{}, titles[], states[], wage_median, dependent_share,
    naics, fy}} for certified, full-time H-1B filings in her role SOC codes."""
    agg = {}
    for p_ in paths:
        rows = iter_xlsx_rows(p_) if str(p_).endswith(".xlsx") else csv.DictReader(open(p_, newline="", encoding="utf-8", errors="replace"))
        for r in rows:
            if not str(r.get("CASE_STATUS", "")).lower().startswith("certified"):
                continue
            if str(r.get("VISA_CLASS", "H-1B")).upper() not in ("H-1B", ""):
                continue
            soc = str(r.get("SOC_CODE", ""))[:7]
            g = SOC_GROUP.get(soc)
            if not g or str(r.get("FULL_TIME_POSITION", "Y")).upper().startswith("N"):
                continue
            name = str(r.get("EMPLOYER_NAME", "")).strip()
            k = norm_company(name)
            if not k:
                continue
            a = agg.setdefault(k, {"name": name, "total": 0, "groups": Counter(), "titles": Counter(), "states": Counter(),
                                   "wages": [], "dep": 0, "naics": str(r.get("NAICS_CODE", ""))[:4], "fy": ""})
            a["total"] += 1
            a["groups"][g] += 1
            a["titles"][str(r.get("JOB_TITLE", "")).strip().title()[:60]] += 1
            a["states"][str(r.get("WORKSITE_STATE", "")).strip()] += 1
            w_ = annual_wage(r.get("WAGE_RATE_OF_PAY_FROM"), r.get("WAGE_UNIT_OF_PAY", "Year"))
            if w_:
                a["wages"].append(w_)
            if str(r.get("H-1B_DEPENDENT", r.get("H1B_DEPENDENT", ""))).upper().startswith("Y"):
                a["dep"] += 1
            fy = re.search(r"FY(\d{4})", Path(p_).name)
            a["fy"] = fy.group(1) if fy else a["fy"]
    out = {}
    for k, a in agg.items():
        ws = sorted(a["wages"])
        out[k] = {"name": a["name"], "total": a["total"], "groups": dict(a["groups"]), "titles": [t for t, _ in a["titles"].most_common(5)],
                  "states": [s for s, _ in a["states"].most_common(4)], "wage_median": round(ws[len(ws) // 2]) if ws else None,
                  "dependent_share": round(a["dep"] / a["total"], 2), "naics": a["naics"], "fy": a["fy"]}
    return out


def download_lca(outdir, max_files=2):
    """Newest published LCA disclosure file of the current fiscal year (cumulative) plus the previous year's Q4."""
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    fy_now = NOW.year + (1 if NOW.month >= 10 else 0)
    got = []
    for fy in (fy_now, fy_now - 1, fy_now - 2):
        for q in (4, 3, 2, 1):
            if len(got) >= max_files or (got and fy == int(re.search(r"FY(\d{4})", got[-1].name).group(1))):
                break
            url = SP_CFG["url_template"].format(fy=fy, q=q)
            dest = outdir / f"LCA_FY{fy}_Q{q}.xlsx"
            try:
                req = urllib.request.Request(url, headers={"User-Agent": UA})
                with urllib.request.urlopen(req, timeout=600) as r, open(dest, "wb") as f:
                    while True:
                        chunk = r.read(1 << 20)
                        if not chunk:
                            break
                        f.write(chunk)
                got.append(dest)
                log(f"downloaded {url} ({dest.stat().st_size / 1e6:.0f} MB)")
                break
            except Exception as e:
                log(f"not available: {url} ({type(e).__name__})")
    return got


_SPONSORS = None


def load_sponsors(path=None):
    global _SPONSORS
    if _SPONSORS is None or path:
        p_ = Path(path or os.environ.get("JOBPIPE_SPONSORS", HERE / "sponsors.json"))
        _SPONSORS = load_json(p_, default={}) if p_.exists() else {}
    return _SPONSORS


def sponsor_lookup(company):
    db = load_sponsors()
    if not db:
        return None
    k = norm_company(company)
    if k in db:
        return db[k]
    cands = [x for x in db if x.startswith(k + " ") or k.startswith(x + " ")]
    return db[max(cands, key=len)] if cands else None


def grow_candidates(sponsors, rows):
    g = SP_CFG.get("grow", {})
    have = {norm_company(r["company"]) for r in rows} | {norm_company(a) for r in rows for a in (r.get("aliases") or "").split(";") if a.strip()}
    out = []
    for k, s_ in sponsors.items():
        if k in have:
            continue
        gs = s_["groups"]
        target = gs.get("data_science", 0) + gs.get("analytics", 0) + gs.get("finance_risk", 0) + gs.get("software", 0)
        core = gs.get("data_science", 0) + gs.get("finance_risk", 0)
        if target < g.get("min_lca_target_socs", 10) and core < g.get("min_lca_data_finance", 3):
            continue
        if s_["dependent_share"] >= g.get("exclude_h1b_dependent_share", 0.5):
            continue
        low = s_["name"].lower()
        if any(d in low for d in GATES["defense_employers"]) or any(a in low for a in GATES["staffing_agencies"]) \
                or re.search(r"universit|college|school|hospital|health system|medical center|state of |county|city of", low):
            continue
        out.append((core * 3 + target, k, s_))
    out.sort(key=lambda x: -x[0])
    return out


def do_registry_grow(sponsors, rows, limit=None, probe_boards=True):
    """Add top DOL sponsors whose job board answers (cheap to harvest). Sponsors without a findable board go
    to registry_candidates.csv instead of the daily career-page ledger (no browser cost until a human or a
    discovery find gives a careers URL)."""
    limit = limit or SP_CFG.get("grow", {}).get("max_new_per_run", 200)
    added, pending = [], []
    for score, k, s_ in grow_candidates(sponsors, rows):
        if len(added) >= limit:
            break
        found = None
        if probe_boards:
            for a in ("greenhouse", "ashby", "lever"):
                for slug_ in slug_candidates(s_["name"]):
                    ok, n = probe(a, slug_)
                    if ok and n > 0:
                        found = (a, slug_)
                        break
                if found:
                    break
        sector = CONFIG["sectors"].get("naics_hint", {}).get(s_.get("naics", ""), "")
        base = {k2: "" for k2 in CSV_FIELDS} | {"company": s_["name"].title(), "sector": sector or ("finance" if s_["groups"].get("finance_risk") else "tech"),
                                                  "startup": "no", "atlanta": "yes" if "GA" in s_["states"] else "no", "source": "dol",
                                                  "added": TODAY.isoformat(), "lca_target": str(sum(s_["groups"].values())), "lca_fy": s_["fy"]}
        if found:
            base.update(ats=found[0], ats_slug=found[1], status="check",
                        notes=f"check: added from DOL LCA data (FY{s_['fy']}: {sum(s_['groups'].values())} filings in her role codes; e.g. {', '.join(s_['titles'][:2])}); board guessed by name, confirm it is this employer")
            added.append(base)
        else:
            pending.append(base | {"notes": f"DOL sponsor, no board found yet; titles: {', '.join(s_['titles'][:3])}"})
    return added, pending


def registry_additions_from_rows(new_rows, rows):
    """Companies of rows just added to a tracker that the registry doesn't know: discovered companies."""
    have = {norm_company(r["company"]) for r in rows}
    out = {}
    for r in new_rows:
        k = norm_company(r.get("company"))
        if not k or k in have or k in out:
            continue
        link = r.get("apply_link") or ""
        ats, slug_, wurl = "", "", ""
        for a, rx in LINK_ATS:
            m = rx.search(link)
            if m:
                if a == "workday":
                    ats, wurl = "workday", f"https://{m.group(1)}.{m.group(2)}.myworkdayjobs.com/{m.group(3)}"
                elif a == "oracle":
                    ats, slug_ = "oracle", f"{m.group(1).lower()}|{m.group(2)}"
                else:
                    ats, slug_ = a, m.group(1)
                break
        pu = urllib.parse.urlparse(link)
        out[k] = {k2: "" for k2 in CSV_FIELDS} | {"company": r["company"], "sector": r.get("sector", ""), "startup": "no",
                                                   "atlanta": "yes" if "atlanta" in (r.get("location") or "").lower() else "no",
                                                   "careers_url": f"{pu.scheme}://{pu.netloc}/" if pu.netloc else "", "ats": ats or "custom",
                                                   "ats_slug": slug_, "workday_url": wurl, "status": "check", "source": "discovered",
                                                   "added": TODAY.isoformat(),
                                                   "notes": f"check: discovered {TODAY.isoformat()} via {r.get('source', '?')} ({r.get('title', '')[:50]}, score {r.get('score')})"}
    return list(out.values())

# =========================================================================== registry: backfill and check
LINK_ATS = [("workday", WD_RE), ("oracle", re.compile(r"https?://([a-z0-9.-]+\.oraclecloud\.com)/hcmUI/CandidateExperience/[a-z_-]+/sites/([A-Za-z0-9_]+)", re.I)),
            ("greenhouse_eu", re.compile(r"job-boards\.eu\.greenhouse\.io/([a-z0-9_-]+)/jobs", re.I)),
            ("greenhouse", re.compile(r"(?:job-boards|boards)\.greenhouse\.io/(?!embed)([a-z0-9_-]+)/jobs", re.I)),
            ("lever", re.compile(r"jobs\.lever\.co/([a-z0-9_.-]+)/", re.I)),
            ("ashby", re.compile(r"jobs\.ashbyhq\.com/([a-z0-9_.%-]+)/", re.I)),
            ("smartrecruiters", re.compile(r"jobs\.smartrecruiters\.com/([A-Za-z0-9_-]+)/", re.I))]
WD_BAD_SITE = re.compile(r"intern|campus|contractor|restricted|private|universit|student|bottler|internal", re.I)


def links_from_files(paths):
    links = []
    for p in paths:
        p = str(p)
        if p.endswith(".md"):
            links += [(j["company"], j["url"]) for j in parse_speedyapply(Path(p).read_text())]
            continue
        data = load_json(p)
        for r in data if isinstance(data, list) else []:
            links.append((r.get("company") or r.get("company_name", ""), r.get("apply_link") or r.get("url", "")))
    return [(c, u) for c, u in links if c and u]


GH_BAD_SLUG = re.compile(r"private|internal|test|sandbox|alt", re.I)
CAMPUS = re.compile(r"campus|universit|student|entry|graduate|early", re.I)


def do_backfill(links, rows=None, write=True):
    """Fill boards only for companies that have none. A company that already has a board is never
    overwritten: a different observed board is written to its notes ('observed: ats/slug') for
    registry-check and a human to confirm (a same-name company can be a different employer,
    e.g. Figure vs Figure AI). Up to two boards are kept (main + campus)."""
    rows = rows if rows is not None else read_companies_csv()
    by = defaultdict(list)
    for c, u in links:
        by[norm_company(c)].append(u)
    changes = []
    for r in rows:
        urls = by.get(norm_company(r["company"]), []) + [u for a in (r.get("aliases") or "").split(";") if a.strip()
                                                          for u in by.get(norm_company(a), [])]
        if not urls:
            continue
        found = None
        for ats, rx in LINK_ATS:
            hits = [m for u in urls for m in [rx.search(u)] if m]
            if not hits:
                continue
            if ats == "workday":
                vals = Counter(f"https://{m.group(1)}.{m.group(2)}.myworkdayjobs.com/{m.group(3)}" for m in hits)
                bad = lambda v: WD_BAD_SITE.search(v.rsplit("/", 1)[1]) and not CAMPUS.search(v.rsplit("/", 1)[1])
            elif ats == "oracle":
                vals = Counter(f"{m.group(1).lower()}|{m.group(2)}" for m in hits)
                bad = lambda v: False
            else:
                vals = Counter(m.group(1) for m in hits)
                bad = lambda v: bool(GH_BAD_SLUG.search(v))
            good = [v for v, _ in vals.most_common() if not bad(v)]
            if not good:
                break
            main = [v for v in good if not CAMPUS.search(v.rsplit("/", 1)[-1])]
            campus = [v for v in good if CAMPUS.search(v.rsplit("/", 1)[-1])]
            picked = (main[:1] + campus[:1]) or good[:1]
            found = (ats, ";".join(picked), f"{len(vals)} board(s) seen")
            break
        if not found:
            continue
        ats, val, note = found
        key = "workday_url" if ats == "workday" else "ats_slug"
        human = (r.get("notes") or "").split(" | check: ")[0].strip()
        human = "" if human.startswith("check: ") else human
        if has_feed(r):
            cur = r.get("workday_url") if (r.get("ats") or "") == "workday" else f"{r.get('ats')}/{r.get('ats_slug')}"
            obs = val if ats == "workday" else f"{ats}/{val}"
            if obs != cur and "observed: " + obs not in (r.get("notes") or ""):
                machine = (r.get("notes") or "").split(" | check: ")[1] if " | check: " in (r.get("notes") or "") else ""
                machine = "; ".join(x for x in [machine, f"observed: {obs}"] if x)
                r["notes"] = (human + " | check: " + machine) if human else "check: " + machine
                changes.append({"company": r["company"], "old": cur, "new": f"(kept) observed {obs}"})
            continue
        changes.append({"company": r["company"], "old": f"{r.get('ats')}/{r.get('ats_slug') or r.get('workday_url')}", "new": f"{ats}/{val}"})
        r["ats"] = ats
        r[key] = val
        if key == "workday_url":
            r["ats_slug"] = ""
        r["status"] = "check"
        msg = f"backfilled {TODAY.isoformat()} from observed posting links ({note}); registry-check confirms"
        r["notes"] = (human + " | check: " + msg) if human else ("check: " + msg)
    if write:
        write_companies_csv(rows)
    return changes


ATS_PATTERNS = [("workday", WD_RE)] + [(a, rx) for a, rx in LINK_ATS if a != "workday"] + [
    ("greenhouse", re.compile(r"boards-api\.greenhouse\.io/v1/boards/([a-z0-9_-]+)", re.I)),
    ("greenhouse", re.compile(r"greenhouse\.io/embed/job_board\?for=([a-z0-9_-]+)", re.I))]
JUNK_SLUGS = {"embed", "js", "api", "v1", "boards", "jobs", "job_board", "static", "job_app"}


def slug_candidates(name):
    base = re.sub(r"\(.*?\)", "", name).lower().replace("&", "and")
    words = [w for w in re.findall(r"[a-z0-9]+", base) if w not in ("the", "inc", "llc", "company", "group")]
    cands = ["".join(words), "-".join(words)]
    if words:
        cands += [words[0], words[0] + "hq", words[0] + "inc", "".join(words) + "jobs", "".join(words) + "careers"]
    return list(dict.fromkeys(c for c in cands if c))


def probe(ats, slug_):
    url = {"greenhouse": f"https://boards-api.greenhouse.io/v1/boards/{slug_}/jobs",
           "greenhouse_eu": f"https://boards-api.eu.greenhouse.io/v1/boards/{slug_}/jobs",
           "lever": f"https://api.lever.co/v0/postings/{slug_}?mode=json&limit=1",
           "ashby": f"https://api.ashbyhq.com/posting-api/job-board/{slug_}",
           "smartrecruiters": f"https://api.smartrecruiters.com/v1/companies/{slug_}/postings?limit=1"}[ats]
    st, body, _ = http(url, retries=0, timeout=15)
    if st != 200:
        return False, 0
    try:
        d = json.loads(body)
    except Exception:
        return False, 0
    n = len(d.get("jobs", [])) if isinstance(d, dict) and "jobs" in d else (d.get("totalFound", 0) if isinstance(d, dict) else len(d))
    if ats == "smartrecruiters" and n == 0:
        return False, 0
    return True, n


def keep_note(c, notes):
    """Keep the human note (before ' | check: ') and any 'observed: …' boards; replace the rest."""
    human = (c.get("notes") or "").split(" | check: ")[0]
    human = "" if human.startswith("check: ") else human.strip()
    obs = [f"{k}: {o}" for k, o in re.findall(r"(observed|guess): ([^;\s]+)", c.get("notes") or "")]
    machine = "; ".join(list(dict.fromkeys(notes + obs)))
    return (human + (" | check: " + machine if machine else "")) if human else ("check: " + machine if machine else "")


def check_company(c):
    """Probe every board a company has; if none answers, try boards observed in notes, then the
    career page's links, then guessed slugs (guesses are marked 'check' for a human)."""
    c = dict(c)
    notes = []
    ats = (c.get("ats") or "").lower()
    if ats in ("greenhouse", "greenhouse_eu", "lever", "ashby", "smartrecruiters") and c.get("ats_slug"):
        good = []
        for s_ in [x.strip() for x in c["ats_slug"].split(";") if x.strip()]:
            ok, n = probe(ats, s_)
            (good if ok else notes).append(f"{ats}/{s_}: {n} jobs" if ok else f"{ats}/{s_} failed")
            if ok:
                c.setdefault("_ok_slugs", []).append(s_)
        if good:
            c["ats_slug"] = ";".join(c.pop("_ok_slugs"))
            c.update(status="ok", last_checked=TODAY.isoformat(), notes=keep_note(c, good + notes))
            return c, None
    if ats == "workday" and c.get("workday_url"):
        good = []
        for u in [x.strip() for x in c["workday_url"].split(";") if x.strip()]:
            st, jobs = fetch_workday({**c, "workday_url": u}, queries=["data"])
            if st == 200:
                good.append(u)
            else:
                notes.append(f"workday {u.rsplit('/', 1)[-1]} failed ({st})")
        if good:
            c.update(workday_url=";".join(good), status="ok", last_checked=TODAY.isoformat(), notes=keep_note(c, notes))
            return c, None
    if ats == "oracle" and c.get("ats_slug"):
        st, jobs = fetch_oracle(c, queries=["data"])
        if st == 200:
            c.update(status="ok", last_checked=TODAY.isoformat(), notes=keep_note(c, [f"oracle: {len(jobs)} jobs for 'data'"]))
            return c, None
        notes.append(f"oracle {c['ats_slug']} failed ({st})")
    if ats == "amazon":
        st, _ = fetch_amazon(c, queries=["data scientist"])
        c.update(status="ok" if st == 200 else "broken", last_checked=TODAY.isoformat())
        return c, None
    found = None
    for obs in re.findall(r"(?:observed|guess): ([^;\s]+)", c.get("notes") or ""):
        if "myworkdayjobs.com" in obs:
            st, _ = fetch_workday({**c, "workday_url": obs}, queries=["data"])
            if st == 200:
                found = ("workday", "", obs)
        else:
            a, _, s_ = obs.partition("/")
            if a in ("greenhouse", "greenhouse_eu", "lever", "ashby", "smartrecruiters") and probe(a, s_)[0]:
                found = (a, s_, "")
        if found:
            notes.append(f"switched to observed board {obs}")
            break
    if not found and c.get("careers_url"):
        st, body, _ = http(c["careers_url"], headers={"Accept": "text/html"}, retries=1)
        if st == 200:
            for a, rx in ATS_PATTERNS:
                m = rx.search(body)
                if m:
                    if a == "workday":
                        found = ("workday", "", m.group(0))
                    elif a == "oracle":
                        found = ("oracle", f"{m.group(1).lower()}|{m.group(2)}", "")
                    elif m.group(1).strip("/").lower() not in JUNK_SLUGS:
                        found = (a, m.group(1).strip("/").lower(), "")
                    if found:
                        break
        else:
            notes.append(f"careers page HTTP {st}")
        c["_careers_http"] = st
    if not found:
        for a in ("greenhouse", "ashby", "lever"):
            for s_ in slug_candidates(c["company"]):
                ok, n = probe(a, s_)
                if ok and n > 0:
                    found = (a, s_, "")
                    notes.append(f"guessed {a}/{s_} ({n} jobs) — confirm it is the right company")
                    break
            if found:
                break
    change = None
    if found:
        a, s_, wurl = found
        change = {"company": c["company"], "old": f"{c.get('ats')}/{c.get('ats_slug') or c.get('workday_url')}", "new": f"{a}/{s_ or wurl}"}
        c["ats"], c["ats_slug"] = a, s_
        if wurl:
            c["workday_url"] = wurl
        c["status"] = "check" if any("guessed" in n for n in notes) else "ok"
    else:
        c["status"] = "no_feed" if c.get("_careers_http") == 200 else "broken"
    c.pop("_careers_http", None)
    c["last_checked"] = TODAY.isoformat()
    c["notes"] = keep_note(c, notes)
    return c, change


def do_registry_check(sectors=None, workers=6, write=True):
    allrows = read_companies_csv()
    target = [r for r in allrows if not sectors or company_sectors(r) & set(sectors)]
    with cf.ThreadPoolExecutor(max_workers=workers) as ex:
        results = list(ex.map(check_company, target))
    upd = {r["company"]: r for r, _ in results}
    new_rows = [upd.get(r["company"], r) for r in allrows]
    if write:
        write_companies_csv(new_rows)
    return Counter(r["status"] for r, _ in results), [ch for _, ch in results if ch]


# =========================================================================== legacy v1 helpers (kept for old prompts)
def do_newgrad(days, tracker_rows, path=None):
    st, jobs = fetch_simplify(path)
    if st != 200:
        raise SystemExit(f"could not download SimplifyJobs list (HTTP {st})")
    scr = [j for j in do_screen(jobs, days * 24) if j["screen"] == "pass" and j["tracker"] == "newgrad"]
    out = [j for j in do_dedupe(scr, tracker_rows) if j["screen"] == "pass"]
    for j in out:
        j["quick_fit"] = band(j.get("pre_score", 0)).replace("Stretch", "Low").replace("Below", "Low")
    return sorted(out, key=lambda j: j.get("posted_at", ""), reverse=True), funnel(scr)


def route_one(j):
    prog, fam = classify(j)
    j["tracker"] = "newgrad" if prog == "newgrad" else "main"
    return j


# =========================================================================== CLI
def table(rows, cols):
    w = {c: max([len(c)] + [len(str(r.get(c, ""))) for r in rows]) for c in cols}
    print(" | ".join(c.ljust(w[c]) for c in cols))
    print("-+-".join("-" * w[c] for c in cols))
    for r in rows:
        print(" | ".join(str(r.get(c, "")).ljust(w[c]) for c in cols))


def compact(j, extra_age=0.0):
    a = j.get("age_hours")
    age = f"{a + extra_age:.0f}h" if a is not None else "?"
    return (f"{j.get('pre_score', '-'):>3}|{age}|{short(j.get('company'), 26)}|{short(j.get('title'), 58)}|"
            f"{short(j.get('location'), 24)}|{j.get('tracker', '')}|{j.get('source', '')}|"
            f"needs: {','.join(j.get('needs', [])) or '-'}|{short('; '.join(j.get('flags', [])), 60)}")


def main(argv=None):
    ap = argparse.ArgumentParser(prog="jobpipe", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("doctor", help="what this machine can reach")
    p = sub.add_parser("domains", help="hosts to allow in claude.ai Settings > Capabilities (fetch option A)")
    p = sub.add_parser("companies"); p.add_argument("--sector"); p.add_argument("--fair", action="store_true"); p.add_argument("--no-feed", action="store_true")
    p = sub.add_parser("collect", help="every machine-readable source -> raw.json [network]")
    p.add_argument("--out", default="raw.json"); p.add_argument("--hours", type=float, default=72); p.add_argument("--only", nargs="*")
    p.add_argument("--no-enrich", action="store_true"); p.add_argument("--simplify-file"); p.add_argument("--speedy-files", nargs="*")
    p = sub.add_parser("plan", help="build the ledger for today's run")
    p.add_argument("--out", default="ledger.json"); p.add_argument("--hours", type=float); p.add_argument("--sites", nargs="*")
    p.add_argument("--skip", nargs="*"); p.add_argument("--no-career-pages", action="store_true")
    p.add_argument("--raw", help="raw.json from collect: marks the API items done / blocked")
    p = sub.add_parser("ledger", help="show the ledger: summary, or the next items to do")
    p.add_argument("file"); p.add_argument("--next", type=int, default=0); p.add_argument("--site")
    p = sub.add_parser("mark", help="mark ledger item(s) done or blocked")
    p.add_argument("file"); p.add_argument("--item", required=True); p.add_argument("--status", choices=["done", "blocked", "todo"], required=True)
    p.add_argument("--results", type=int); p.add_argument("--kept", type=int); p.add_argument("--blocker", default=""); p.add_argument("--note", default="")
    p = sub.add_parser("ingest", help="cards Claude read from a site -> pool.jsonl; marks the ledger item done")
    p.add_argument("file", help="ledger.json"); p.add_argument("--item", required=True); p.add_argument("--cards", required=True, help="file with one card per line ('-' = stdin)")
    p.add_argument("--pool", default="pool.jsonl"); p.add_argument("--company")
    p = sub.add_parser("screen", help="gates + routing + pre-score -> screened.json")
    p.add_argument("inputs", nargs="+", help="raw.json and/or pool.jsonl"); p.add_argument("--hours", type=float)
    p.add_argument("--seen", help="state.json (first-seen memory; updated in place)"); p.add_argument("--h1b"); p.add_argument("--out", default="screened.json")
    p = sub.add_parser("shortlist", help="dedupe vs trackers, resolve to employer postings -> shortlist.json")
    p.add_argument("screened"); p.add_argument("--tracker", action="append", default=[]); p.add_argument("--raw", help="raw.json (feed jobs for resolving)")
    p.add_argument("--threshold", type=float); p.add_argument("--out", default="shortlist.json"); p.add_argument("--print", type=int, default=60)
    p.add_argument("--skip-fps", default="scored_fps.txt", help="fingerprints already handled today (written by finalize)")
    p = sub.add_parser("brief", help="requirement lines of shortlisted postings, for scoring with few tokens")
    p.add_argument("shortlist"); p.add_argument("--start", type=int, default=0); p.add_argument("--n", type=int, default=8)
    p.add_argument("--fp", nargs="*"); p.add_argument("--chars", type=int, default=1100)
    p = sub.add_parser("attach", help="after opening a posting: store its text / employer link / date and re-run the gates")
    p.add_argument("shortlist"); p.add_argument("--fp", required=True); p.add_argument("--text", help="file with the posting text ('-' = stdin)")
    p.add_argument("--apply-url"); p.add_argument("--posted", help="ISO date or a relative age such as 'Yesterday'"); p.add_argument("--req-id")
    p.add_argument("--company"); p.add_argument("--location"); p.add_argument("--hours", type=float)
    p = sub.add_parser("finalize", help="Claude's scores -> validated rows, merged trackers")
    p.add_argument("scored", help="shortlist.json (with --scores) or a list that already carries score blocks")
    p.add_argument("--scores", help="scores.jsonl written by Claude, merged into the shortlist by fp"); p.add_argument("--main", required=True); p.add_argument("--newgrad", required=True)
    p.add_argument("--out-main", default="roles_main_next.json"); p.add_argument("--out-newgrad", default="roles_newgrad_next.json")
    p.add_argument("--summary", default="finalize_summary.json"); p.add_argument("--done-fps", default="scored_fps.txt")
    p = sub.add_parser("audit", help="ledger complete? (legacy: --run R source_log.json)")
    p.add_argument("file"); p.add_argument("--run")
    p = sub.add_parser("report"); p.add_argument("--ledger", required=True); p.add_argument("--screened"); p.add_argument("--summary"); p.add_argument("--raw"); p.add_argument("--out", default="report.md")
    p = sub.add_parser("tracker-docs", help="rows -> per-document files + ArtifactData batch manifests (50 per batch)")
    p.add_argument("rows"); p.add_argument("--existing", help="tracker.json exported from the same artifact (ids already there)"); p.add_argument("--outdir", default="docs_out")
    p = sub.add_parser("tracker-export", help="ArtifactData list out_dir -> tracker.json")
    p.add_argument("folder"); p.add_argument("--out", default="tracker.json")
    p = sub.add_parser("referrals", help="referral packs (links + draft), one per company with a High open role, create-only")
    p.add_argument("tracker"); p.add_argument("--packs", help="folder from ArtifactData list of `referrals` (out_dir); those companies are skipped")
    p.add_argument("--outdir", default="referral_out")
    p = sub.add_parser("validate"); p.add_argument("rows"); p.add_argument("--kind", choices=["main", "newgrad"], default="main")
    p = sub.add_parser("stats"); p.add_argument("tracker")
    p = sub.add_parser("calibrate", help="does the score predict what she applies to?"); p.add_argument("tracker")
    p = sub.add_parser("registry-backfill", help="fill missing boards from observed posting links"); p.add_argument("links", nargs="+"); p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("registry-check", help="verify boards, detect ATS from career pages [network]"); p.add_argument("--sector"); p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("sponsors-build", help="DOL LCA disclosure file(s) -> sponsors.json [network with --download]")
    p.add_argument("files", nargs="*"); p.add_argument("--download", metavar="DIR"); p.add_argument("--out", default=str(HERE / "sponsors.json"))
    p = sub.add_parser("registry-grow", help="add top DOL sponsors whose job board answers [network unless --no-probe]")
    p.add_argument("--sponsors", default=str(HERE / "sponsors.json")); p.add_argument("--limit", type=int); p.add_argument("--no-probe", action="store_true")
    p.add_argument("--dry-run", action="store_true"); p.add_argument("--candidates", default="registry_candidates.csv")
    p = sub.add_parser("registry-add", help="merge registry_additions.csv (companies discovered by the daily run) into companies.csv")
    p.add_argument("additions"); p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("h1b-build", help="USCIS H-1B Employer Data Hub CSV(s) -> h1b.json"); p.add_argument("csvs", nargs="+"); p.add_argument("--out", default=str(HERE / "h1b.json"))
    p = sub.add_parser("h1b", help="look up an employer"); p.add_argument("company"); p.add_argument("--db")
    p = sub.add_parser("state", help="first-seen memory: new | info | unpack (runstate docs -> state.json) | pack (this run's new entries -> one doc)")
    p.add_argument("action", choices=["new", "info", "unpack", "pack"]); p.add_argument("file", help="state.json (unpack: the out_dir of an ArtifactData list of `runstate`)")
    p.add_argument("--out", help="unpack: state.json to write; pack: folder for the doc file and batch manifest")
    p.add_argument("--since", help="pack: ISO time the run started (entries first seen since then are this run's)")
    # legacy
    p = sub.add_parser("fetch"); p.add_argument("--run", default="sweep"); p.add_argument("--out", default="raw_jobs.json")
    p = sub.add_parser("dedupe"); p.add_argument("candidates"); p.add_argument("--tracker", action="append", default=[]); p.add_argument("--out", default="deduped.json")
    p.add_argument("--li-days", type=float, default=3); p.add_argument("--li-ref", type=int)
    p = sub.add_parser("route"); p.add_argument("candidates"); p.add_argument("--out", default="routed.json")
    p = sub.add_parser("newgrad"); p.add_argument("--days", type=int, default=7); p.add_argument("--tracker", action="append", default=[]); p.add_argument("--listings"); p.add_argument("--out", default="newgrad_candidates.json")
    a = ap.parse_args(argv)

    if a.cmd == "doctor":
        mode, res = doctor()
        for k, v in res.items():
            print(f"{k:32} HTTP {v}")
        print(f"\nMODE: {mode}")
        print({"direct": "Job boards are reachable: run `collect` here.",
               "github_only": "Job boards are blocked but GitHub works: download data/latest/raw.json from the feed repo "
                              "(references/automation.md) or, if none, run `collect --only simplify speedyapply` and read "
                              "employer feeds with WebFetch.",
               "none": "No network for sources: use WebFetch for feeds and say so in the report."}[mode])
    elif a.cmd == "domains":
        hosts = {"boards-api.greenhouse.io", "boards-api.eu.greenhouse.io", "api.lever.co", "api.ashbyhq.com",
                 "api.smartrecruiters.com", "www.amazon.jobs", "hn.algolia.com", "raw.githubusercontent.com"}
        for c in read_companies_csv():
            for u in (c.get("workday_url") or "").split(";"):
                m = WD_RE.search(u)
                if m:
                    hosts.add(f"{m.group(1)}.{m.group(2)}.myworkdayjobs.com")
            if c.get("ats") == "oracle" and "|" in (c.get("ats_slug") or ""):
                hosts.add(c["ats_slug"].split("|")[0])
        wild = ["*.myworkdayjobs.com", "*.oraclecloud.com"]
        print("If wildcards are accepted, these cover everything:")
        print("\n".join(sorted(h for h in hosts if "myworkdayjobs" not in h and "oraclecloud" not in h) + wild))
        print(f"\nOtherwise, all {len(hosts)} hosts:")
        print("\n".join(sorted(hosts)))
    elif a.cmd == "companies":
        rows = load_companies([a.sector] if a.sector else None, a.fair)
        if a.no_feed:
            rows = [r for r in rows if not has_feed(r)]
        table(rows, ["company", "sector", "startup", "gt_fair", "fair_sponsor", "ats", "ats_slug", "status"])
        print(f"\n{len(rows)} companies; with a feed: {sum(1 for r in rows if has_feed(r))}")
    elif a.cmd == "collect":
        raw = do_collect(a.hours, enrich=not a.no_enrich, only=a.only,
                         offline={"simplify": a.simplify_file, "speedyapply": a.speedy_files})
        dump_json(raw, a.out)
        print(json.dumps(raw["meta"]["sources"], indent=1))
        print(f"{len(raw['jobs'])} jobs -> {a.out}; feeds failed: {len(raw['meta']['feeds_failed'])}; no feed: {len(raw['meta']['no_feed'])}")
    elif a.cmd == "plan":
        led = build_plan(a.hours, a.sites, a.skip, career_pages=not a.no_career_pages)
        if a.raw:
            apply_collect_meta(led, load_json(a.raw, default={}).get("meta"))
        dump_json(led, a.out)
        s = ledger_summary(led)
        table([{"site": k, "items": sum(v[x] for x in ("todo", "done", "blocked"))} for k, v in s.items()], ["site", "items"])
        br = [it for it in led["items"] if it["mode"] == "browser"]
        print(f"\n{len(led['items'])} ledger items ({len(br)} browser pages, ~{est_minutes(br)} min of browsing) -> {a.out}")
    elif a.cmd == "ledger":
        led = load_json(a.file, default={})
        if a.next:
            todo = [it for it in led["items"] if it["status"] == "todo" and (not a.site or it["site"] == a.site)]
            todo.sort(key=lambda x: x.get("priority", 9))
            for it in todo[: a.next]:
                print(f"{it['id']}\n  {it['url'] or '(api)'}\n  q: {it['query']}")
            print(f"\n{len(todo)} items left")
        else:
            ok, rows = do_audit_ledger(led)
            table(rows, ["site", "done", "blocked", "todo", "results", "status"])
    elif a.cmd == "mark":
        led = load_json(a.file, default={})
        if a.status == "blocked" and not valid_blocker(a.blocker):
            print("WARNING: a blocker counts only when it quotes a security check, CAPTCHA, sign-in wall, HTTP 429 or "
                  "'computer unreachable'. Friction (empty list, slow page) is a bug to fix, not a blocker.")
        hit = mark_item(led, a.item, a.status, a.results, a.kept, a.blocker, a.note)
        dump_json(led, a.file)
        print(f"marked {len(hit)} item(s) {a.status}")
    elif a.cmd == "ingest":
        led = load_json(a.file, default={})
        text = sys.stdin.read() if a.cards == "-" else Path(a.cards).read_text(encoding="utf-8")
        rows = do_ingest(led, a.item, text, a.pool, extra={"company": a.company} if a.company else None)
        dump_json(led, a.file)
        flagged = sum(1 for r in rows if r.get("reposted"))
        known = sum(1 for r in rows if r.get("known_company"))
        print(f"{len(rows)} cards from {a.item} -> {a.pool} ({known} at companies already harvested, dropped; {len(rows) - known} new-company leads; {flagged} reposted); "
              f"no company parsed: {sum(1 for r in rows if not r['company'])}; no age: {sum(1 for r in rows if not r['age_text'])}")
        for r in rows[:5]:
            print(f"  {short(r['age_text'], 10)}|{short(r['company'], 24)}|{short(r['title'], 50)}|{short(r['location'], 22)}")
    elif a.cmd == "screen":
        if a.h1b:
            load_h1b(a.h1b)
        jobs = []
        for path in a.inputs:
            d = load_json(path, default=None)
            jobs += d.get("jobs", []) if isinstance(d, dict) else d
        state = load_json(a.seen, default={}) if a.seen else None
        seen = state.setdefault("seen", {}) if state is not None else None
        out = do_screen(jobs, a.hours or CONFIG["window_hours"]["daily"], seen)
        if get_matcher():
            dump_json(get_matcher().stats(), str(Path(a.out).with_name("jd_stats.json")))
        else:
            print("NOTE: no resumes found (resumes/*.md): skills use the keyword fallback")
        if a.seen:
            cutoff = (NOW - dt.timedelta(days=60)).isoformat()
            state["seen"] = {k: v for k, v in seen.items() if v > cutoff}
            state["updated_at"] = NOW.isoformat()
            dump_json(state, a.seen)
        dump_json(out, a.out)
        f = funnel(out)
        print(json.dumps(f))
        by = Counter((j["tracker"], j["screen"]) for j in out)
        print(f"{len(out)} screened -> {a.out}; passed: main {by[('main', 'pass')]}, newgrad {by[('newgrad', 'pass')]}")
    elif a.cmd == "shortlist":
        trk = [r for t in a.tracker for r in load_rows(t)]
        raw = load_json(a.raw, default={}) if a.raw else {}
        feed = [j for j in raw.get("jobs", []) if j.get("source_group") == "Employer job board"] if raw else None
        pool, sl, below = build_shortlist(load_json(a.screened), trk, feed, a.threshold)
        done = set(Path(a.skip_fps).read_text().split()) if a.skip_fps and Path(a.skip_fps).exists() else set()
        if done:
            n0 = len(sl)
            sl = [j for j in sl if j["fp"] not in done]
            print(f"skipped {n0 - len(sl)} roles already handled earlier today ({a.skip_fps})")
        dump_json(sl, a.out)
        dump_json(below, a.out.replace(".json", "_below.json"))
        f = funnel(pool)
        print(json.dumps(f))
        print(f"shortlist {len(sl)} (main {sum(1 for j in sl if j['tracker'] == 'main')}, newgrad {sum(1 for j in sl if j['tracker'] == 'newgrad')}); "
              f"below pre-score threshold {len(below)}; need employer link: {sum(1 for j in sl if 'employer link' in j['needs'])}; "
              f"resolved from feeds: {sum(1 for j in sl if j.get('resolved_from_feed'))} -> {a.out}")
        for j in sl[: a.print]:
            print(compact(j))
    elif a.cmd == "brief":
        sl = load_json(a.shortlist)
        items = [j for j in sl if j["fp"] in set(a.fp)] if a.fp else sl[a.start: a.start + a.n]
        for j in items:
            print(f"## {j['fp']} | {j.get('pre_score')} | {short(j.get('company'), 30)} — {short(j.get('title'), 70)} | "
                  f"{short(j.get('location'), 30)} | {j.get('tracker')} | {j.get('source')}")
            print(f"   link: {j.get('apply_url') or '(none: find the employer posting)'} | posted: {posted_label(j) or '?'} | "
                  f"flags: {short('; '.join(j.get('flags', [])), 120)}")
            print(f"   auto: {j.get('ev_experience_auto') or '-'} | {j.get('h1b_history') or ''}")
            txt = brief_text(j.get("description") or j.get("snippet") or "", a.chars)
            print(f"   req: {txt or '(no posting text: open the posting)'}")
        print(f"\n{len(items)} shown of {len(sl)}; next: --start {a.start + a.n}" if not a.fp else "")
    elif a.cmd == "attach":
        sl = load_json(a.shortlist)
        hit = [i for i, j in enumerate(sl) if j["fp"] == a.fp]
        if not hit:
            raise SystemExit(f"no shortlist item with fp {a.fp}")
        j = dict(sl[hit[0]])
        if a.text:
            j["description"] = (sys.stdin.read() if a.text == "-" else Path(a.text).read_text(encoding="utf-8"))[:8000]
        if a.apply_url:
            j["apply_url"] = a.apply_url
        if a.posted:
            iso_p = iso(a.posted)
            if iso_p:
                j["posted_at_ats"], j["posted_at"] = iso_p, iso_p
                j.pop("age_text", None)
            else:
                j["age_text"], j["posted_at"] = a.posted, ""
        for k in ("req_id", "company", "location"):
            if getattr(a, k):
                j[k] = getattr(a, k)
        keep = {k: j.get(k) for k in ("source", "source_group", "also_on_linkedin", "also_seen_on", "resolved_from_feed")}
        j["exclude_reasons"] = [r for r in j.get("exclude_reasons", []) if r.startswith("duplicate")]
        j["flags"] = [f for f in j.get("flags", []) if not f.startswith("no description")]
        stats_p = Path(a.shortlist).with_name("jd_stats.json")
        res = load_resumes()
        set_matcher(JDMatcher(res, stats=load_json(stats_p, default={})) if res and stats_p.exists() else (JDMatcher(res) if res else None))
        j = screen_one(j, a.hours or CONFIG["window_hours"]["widened"])
        j.update({k: v for k, v in keep.items() if v is not None})
        if j["screen"] == "pass":
            lp = apply_link_problem(j.get("apply_url"))
            if lp:
                j.setdefault("flags", []).append(f"apply link: {lp}")
        j["opened"] = True
        j["needs"] = needs_for(j)
        sl[hit[0]] = j
        dump_json(sl, a.shortlist)
        print(f"{j['fp']}: {j['screen'].upper()} {('; '.join(j['exclude_reasons'])) or ''} | pre-score {j.get('pre_score', '-')} | "
              f"needs: {','.join(j['needs']) or '-'} | {j.get('ev_experience_auto') or ''}")
    elif a.cmd == "finalize":
        main_rows, ng_rows = load_rows(a.main), load_rows(a.newgrad)
        scored = load_json(a.scored)
        if a.scores:
            scored, missing = merge_scores(scored, load_json(a.scores))
            if missing:
                print(f"WARNING: {len(missing)} scores.jsonl lines match no shortlist fp: {missing[:5]}")
        nm, nn, na, leads = do_finalize(scored, main_rows, ng_rows)
        handled = [j["fp"] for j in scored if j.get("score") or j.get("drop") or j.get("screen") == "excluded"]
        if a.done_fps and handled:
            with open(a.done_fps, "a") as f:
                f.write("\n".join(handled) + "\n")
        dump_json(main_rows + nm, a.out_main)
        dump_json(ng_rows + nn, a.out_newgrad)
        adds = registry_additions_from_rows(nm + nn, read_companies_csv())
        if adds:
            ap = Path(a.summary).with_name("registry_additions.csv")
            old_rows = list(csv.DictReader(open(ap, newline="", encoding="utf-8"))) if ap.exists() else []
            seen_c = {norm_company(r["company"]) for r in old_rows}
            with open(ap, "w", newline="", encoding="utf-8") as f:
                w_ = csv.DictWriter(f, fieldnames=CSV_FIELDS, extrasaction="ignore")
                w_.writeheader()
                w_.writerows(old_rows + [r for r in adds if norm_company(r["company"]) not in seen_c])
            print(f"{len(adds)} discovered companies -> {ap} (push to the feed repo: references/automation.md)")
        dump_json({"new_main": nm, "new_newgrad": nn, "not_added": dict(na), "leads": leads, "registry_additions": [r["company"] for r in adds]}, a.summary)
        print(f"main +{len(nm)} -> {a.out_main} ({len(main_rows) + len(nm)} rows); newgrad +{len(nn)} -> {a.out_newgrad} "
              f"({len(ng_rows) + len(nn)} rows); not added: {dict(na)}; leads not added: {len(leads)}")
        for p_ in (a.out_main, a.out_newgrad):
            size = Path(p_).stat().st_size
            if size > 7_000_000:
                print(f"WARNING: {p_} is {size / 1e6:.1f} MB; the dashboard refuses files over 8 MB: archive old rows (tracker.md)")
    elif a.cmd == "audit":
        d = load_json(a.file, default={})
        if isinstance(d, dict) and "items" in d:
            ok, rows = do_audit_ledger(d)
            table(rows, ["site", "done", "blocked", "todo", "results", "status", "blocker"])
        else:
            ok, rows = do_audit_legacy(a.run, d)
            table(rows, ["site", "queries", "status", "blocker"])
        print("\nAUDIT: " + ("PASS" if ok else "INCOMPLETE — finish the todo / SHORT items before scoring or reporting"))
        return 0 if ok else 1
    elif a.cmd == "report":
        led = load_json(a.ledger, default={})
        scr = load_json(a.screened) if a.screened else []
        summ = load_json(a.summary, default={}) if a.summary else {}
        raw = load_json(a.raw, default={}) if a.raw else {}
        txt, ok = build_report(led, funnel(scr), summ.get("new_main", []), summ.get("new_newgrad", []),
                               summ.get("not_added", {}), summ.get("leads", []), raw.get("meta"))
        Path(a.out).write_text(txt, encoding="utf-8")
        print(txt)
    elif a.cmd == "tracker-docs":
        ex = {r.get("_id") or doc_id(r) for r in load_rows(a.existing)} if a.existing else set()
        batches = tracker_docs(load_rows(a.rows), ex, a.outdir)
        n = sum(len(b) for b in batches)
        print(f"{n} new documents in {len(batches)} batch file(s) under {a.outdir}/batch_*.json "
              f"({len(load_rows(a.rows)) - n} skipped: already in the tracker or duplicate ids)")
    elif a.cmd == "tracker-export":
        rows = tracker_export(a.folder)
        dump_json(rows, a.out)
        print(f"{len(rows)} rows -> {a.out}; statuses: {dict(Counter(r.get('status') or '?' for r in rows))}")
    elif a.cmd == "referrals":
        have = {p_.stem for p_ in Path(a.packs).rglob("*.json")} if a.packs else set()
        packs = do_referrals(load_rows(a.tracker), have)
        od = Path(a.outdir); (od / "docs").mkdir(parents=True, exist_ok=True)
        for old in od.glob("batch_*.json"):
            old.unlink()
        writes = []
        for pk in packs:
            i = pk.pop("_id")
            fp_ = od / "docs" / f"{i}.json"
            dump_json(pk, fp_)
            writes.append({"op": "set", "collection": "referrals", "doc_id": i, "file_path": str(fp_.resolve())})
        for n, k in enumerate(range(0, len(writes), 50), 1):
            dump_json(writes[k:k + 50], od / f"batch_{n:03d}.json")
        print(f"{len(packs)} new referral packs ({len(have)} companies already had one) -> {od}/batch_*.json")
    elif a.cmd == "validate":
        probs = validate_rows(load_rows(a.rows), a.kind)
        print("\n".join(probs) if probs else "OK: all rows valid")
        return 1 if probs else 0
    elif a.cmd == "stats":
        print(json.dumps(do_stats(load_rows(a.tracker)), indent=1))
    elif a.cmd == "calibrate":
        print(json.dumps(do_calibrate(load_rows(a.tracker)), indent=1))
    elif a.cmd == "registry-backfill":
        ch = do_backfill(links_from_files(a.links), write=not a.dry_run)
        for c in ch:
            print(f"{c['company']}: {c['old']} -> {c['new']}")
        print(f"\n{len(ch)} companies updated (status 'check' until registry-check confirms)" + (" [dry run]" if a.dry_run else ""))
    elif a.cmd == "registry-check":
        counts, changes = do_registry_check([a.sector] if a.sector else None, write=not a.dry_run)
        print(json.dumps({"status_counts": dict(counts), "changes": changes}, indent=1))
    elif a.cmd == "sponsors-build":
        files = list(a.files) + (download_lca(a.download) if a.download else [])
        if not files:
            raise SystemExit("no LCA files (pass .xlsx/.csv paths or --download DIR)")
        sp = build_sponsors(files)
        dump_json(sp, a.out)
        top = sorted(sp.values(), key=lambda x: -x["total"])[:10]
        print(f"{len(sp)} sponsoring employers in her role codes -> {a.out}; top: " + ", ".join(f"{t['name']} ({t['total']})" for t in top))
    elif a.cmd == "registry-grow":
        rows = read_companies_csv()
        added, pending = do_registry_grow(load_sponsors(a.sponsors), rows, a.limit, probe_boards=not a.no_probe)
        with open(a.candidates, "w", newline="", encoding="utf-8") as f:
            w_ = csv.DictWriter(f, fieldnames=CSV_FIELDS, extrasaction="ignore")
            w_.writeheader()
            w_.writerows(pending)
        if not a.dry_run and added:
            write_companies_csv(rows + added)
        print(f"added {len(added)} sponsors with a working board; {len(pending)} more without a board -> {a.candidates}" + (" [dry run]" if a.dry_run else ""))
    elif a.cmd == "registry-add":
        rows = read_companies_csv()
        have = {norm_company(r["company"]) for r in rows}
        with open(a.additions, newline="", encoding="utf-8") as f:
            new = [r for r in csv.DictReader(f) if norm_company(r.get("company")) not in have]
        if not a.dry_run and new:
            write_companies_csv(rows + [{k: r.get(k, "") for k in CSV_FIELDS} for r in new])
        print(f"{len(new)} discovered companies added" + (" [dry run]" if a.dry_run else "") + ": " + ", ".join(r["company"] for r in new[:20]))
    elif a.cmd == "h1b-build":
        db = {}
        for path in a.csvs:
            fy = (re.search(r"20\d\d", Path(path).name) or [""])[0]
            for k, v in parse_h1b_csv(Path(path).read_text(encoding="utf-8", errors="replace"), fy).items():
                if k not in db or str(v["fy"]) > str(db[k]["fy"]):
                    db[k] = v
        dump_json(db, a.out)
        print(f"{len(db)} employers -> {a.out}")
    elif a.cmd == "h1b":
        if a.db:
            load_h1b(a.db)
        print(json.dumps(h1b_lookup(a.company)))
    elif a.cmd == "state":
        if a.action == "unpack":
            st = state_unpack(a.file)
            dump_json(st, a.out or "state.json")
            print(f"{len(st['seen'])} first-seen entries from {st['docs']} runstate docs -> {a.out or 'state.json'}")
        elif a.action == "pack":
            doc_id_, doc, n = state_pack(load_json(a.file, default={}), a.since or TODAY.isoformat())
            od = Path(a.out or "runstate_out"); od.mkdir(parents=True, exist_ok=True)
            fp_ = od / f"{doc_id_}.json"; dump_json(doc, fp_)
            dump_json([{"op": "set", "collection": "runstate", "doc_id": doc_id_, "file_path": str(fp_.resolve())}], od / "batch_001.json")
            print(f"{n} new first-seen entries -> runstate/{doc_id_} ({fp_.stat().st_size // 1024} KB); manifest {od}/batch_001.json")
        elif a.action == "new":
            dump_json({"seen": {}, "created_at": NOW.isoformat()}, a.file)
            print(f"new state -> {a.file}")
        else:
            st = load_json(a.file, default={})
            print(json.dumps({"seen": len(st.get("seen", {})), "updated_at": st.get("updated_at")}))
    elif a.cmd == "fetch":  # legacy: one sector set
        raw = do_collect(CONFIG["window_hours"]["widened"], only=["feeds"])
        dump_json({"fetched_at": NOW.isoformat(), "feeds": raw["feeds"], "jobs": raw["jobs"]}, a.out)
        print(f"{len(raw['jobs'])} jobs -> {a.out}")
    elif a.cmd == "dedupe":
        trk = [r for t in a.tracker for r in load_rows(t)]
        out = do_dedupe(linkedin_age_check(load_json(a.candidates), a.li_days, a.li_ref), trk)
        dump_json(out, a.out)
        for j in out:
            if j.get("screen") == "excluded":
                print(f"  excluded: {j.get('company')} | {short(j.get('title'), 50)} | {'; '.join(j.get('exclude_reasons', []))}")
        print(f"{sum(1 for j in out if j.get('screen') != 'excluded')} left after dedupe against {len(trk)} tracker rows")
    elif a.cmd == "route":
        out = [route_one(j) for j in load_json(a.candidates)]
        dump_json(out, a.out)
        print(f"main {sum(1 for j in out if j['tracker'] == 'main')}, newgrad {sum(1 for j in out if j['tracker'] == 'newgrad')}")
    elif a.cmd == "newgrad":
        trk = [r for t in a.tracker for r in load_rows(t)]
        out, counts = do_newgrad(a.days, trk, a.listings)
        dump_json(out, a.out)
        print(json.dumps({"kept": len(out), "funnel": counts}, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
